"""Word export of the step 0 characterisation: the .docx must carry exactly what the screen shows."""

import io
import os
import re
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("KEVIDENCE_HOME", str(ROOT))
os.environ.setdefault("OPENAI_API_KEY", "sk-test-not-used")

import characterisation as c  # noqa: E402
import docx_export as dx  # noqa: E402

DEMO = ROOT / "static" / "samples" / "characterisation_demo_dossier.md"
DEMO_META = {"product_name": "DemoVan 99", "product_form": "solid", "routes": ["premixture", "feed"],
             "production_types": ["chemical_synthesis"], "additive_type": "flavouring_compound",
             "components": [{"name": "vanillin", "cas": "121-33-5", "role": "Active substance"}]}


def _load(md):
    return Document(io.BytesIO(dx.markdown_to_docx(md, "Test title")))


@pytest.fixture(scope="module")
def demo_markdown():
    bundle = c.load_domain_bundle("feed")
    doc = c.extract_document(DEMO.name, DEMO.read_bytes(), "D1")
    return c.analyse_dossier([doc], DEMO_META, bundle["primary"]["catalogue"], bundle["primary"]["pages"], datasets=bundle["datasets"])["markdown"]


SAMPLE = """# Title here

*Product form:* solid · *Routes:* feed  
**Draft for review.**

## A heading

Text with **bold**, *italic*, `code`, a pipe \\| and a note[^n1].

| Parameter | Value | Source |
|---|---|---|
| **Specifications** | | |
| assay (%) | 99.8 (99.7–99.9) [3] | D1 p.1 |
| a \\| b | x | D1 p.2 |

Abbreviations: DM: dry matter.
[^n1]: Footnote text with *emphasis*.

- first bullet
  - nested bullet
1. numbered

> **[Scientific officer to conclude]** on compliance.
> second quote line

---

# Annex

Plain.
"""


def test_structure_is_carried_over():
    d = _load(SAMPLE)
    styles = [(p.style.name, p.text) for p in d.paragraphs]
    assert styles[0] == ("Title", "Title here")
    assert ("Heading 2", "A heading") in styles
    assert any(s == "Heading 1" and t == "Annex" for s, t in styles)          # later H1 stays a heading, not a second title
    assert [s for s, _ in styles].count("List Bullet") == 1 and [s for s, _ in styles].count("List Bullet 2") == 1
    assert ("List Number", "numbered") in styles
    assert len(d.tables) == 1


def test_inline_formatting_escapes_and_footnotes():
    d = _load(SAMPLE)
    body = next(p for p in d.paragraphs if p.text.startswith("Text with"))
    runs = {r.text: r for r in body.runs}
    assert runs["bold"].bold and runs["italic"].italic and runs["code"].font.name == "Consolas"
    assert "a pipe | and" in body.text and "\\" not in body.text            # \| unescaped
    marker = next(r for r in body.runs if r.text == "1")
    assert marker.font.superscript
    note = next(p for p in d.paragraphs if p.text.startswith("1 Footnote text"))
    assert note.runs[0].font.superscript and any(r.italic for r in note.runs)
    assert not any("**" in p.text or "[^" in p.text for p in d.paragraphs)


def test_hard_line_break_and_blockquote():
    d = _load(SAMPLE)
    head = next(p for p in d.paragraphs if p.text.startswith("Product form"))
    assert head._p.xpath(".//w:br")                                          # line ending in two spaces
    quote = next(p for p in d.paragraphs if p.text.startswith("[Scientific officer to conclude]"))
    assert quote._p.xpath(".//w:shd") and "second quote line" in quote.text


def test_table_header_repeats_and_cells_match():
    d = _load(SAMPLE)
    t = d.tables[0]
    assert t.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert [c_.text for c_ in t.rows[0].cells] == ["Parameter", "Value", "Source"]
    assert all(r.bold for r in t.rows[0].cells[0].paragraphs[0].runs)
    assert t.rows[3].cells[0].text == "a | b"
    assert t.rows[2].cells[1].text == "99.8 (99.7–99.9) [3]"


def test_page_furniture_and_properties():
    d = _load(SAMPLE)
    sec = d.sections[0]
    assert round(sec.page_width.cm, 1) == 21.0 and round(sec.page_height.cm, 1) == 29.7
    assert "draft for scientific-officer verification" in sec.header.paragraphs[0].text
    footer_xml = sec.footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    assert d.core_properties.title == "Test title" and d.core_properties.author == "KEvidence"


def test_demo_output_roundtrips_every_table_cell_and_heading(demo_markdown):
    d = _load(demo_markdown)
    md_tables = sum(1 for i, ln in enumerate(demo_markdown.splitlines()[:-1])
                    if ln.startswith("|") and re.match(r"^\|?\s*:?-{3,}", demo_markdown.splitlines()[i + 1]))
    assert len(d.tables) == md_tables
    docx_text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c_.text for t in d.tables for r in t.rows for c_ in r.cells)
    headings = [re.sub(r"[*`]", "", ln.lstrip("# ").strip()) for ln in demo_markdown.splitlines() if ln.startswith("#")]
    assert all(h in docx_text for h in headings)
    for ln in demo_markdown.splitlines():
        if ln.startswith("|") and not re.match(r"^\|?\s*:?-{3,}", ln):
            for cell in dx._split_row(ln):
                plain = dx._unescape(re.sub(r"\*\*|`|\[\^[^\]]+\]", "", cell)).replace("*", "").strip()
                if plain:
                    assert plain in docx_text or plain in docx_text.replace("\n", " "), plain
    assert "**" not in docx_text and "[^" not in docx_text


def test_wide_tables_get_content_based_widths():
    rows = [["§", "Requirement", "Missing / to clarify", "Justification located"], ["2.1", "Short", "x " * 80, "—"]]
    widths = dx._column_widths(rows, 4)
    assert widths[2] > widths[0] * 3
    assert widths[3].cm >= 0.2 * len("Justification") + 0.3 - 0.01          # header word never broken
    assert abs(sum(w.cm for w in widths) - dx.TEXT_WIDTH_CM) < 1.5


def test_oversized_markdown_is_refused():
    with pytest.raises(ValueError):
        dx.markdown_to_docx("x" * (dx.MAX_MARKDOWN_CHARS + 1))


# --- HTTP endpoint --------------------------------------------------------------------------

from fastapi.testclient import TestClient  # noqa: E402

import server  # noqa: E402


@pytest.fixture()
def client():
    return TestClient(server.app)


def test_endpoint_returns_a_valid_word_file(client):
    r = client.post("/api/characterisation/export.docx", json={"markdown": SAMPLE, "product_name": "Demo Van/99 ../x"})
    assert r.status_code == 200
    assert r.headers["content-type"] == dx.DOCX_MEDIA_TYPE
    assert r.headers["content-disposition"] == 'attachment; filename="characterisation_Demo_Van_99_._x.docx"'
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert "word/document.xml" in z.namelist()
    assert Document(io.BytesIO(r.content)).tables


@pytest.mark.parametrize("body,status", [({}, 400), ({"markdown": "   "}, 400), ({"markdown": 5}, 400)])
def test_endpoint_rejects_bad_input(client, body, status):
    assert client.post("/api/characterisation/export.docx", json=body).status_code == status


def test_endpoint_rejects_oversized_input(client):
    r = client.post("/api/characterisation/export.docx", json={"markdown": "x" * (dx.MAX_MARKDOWN_CHARS + 1)})
    assert r.status_code == 413


def test_analyse_then_export_equals_the_screen(client):
    with open(DEMO, "rb") as f:
        res = client.post("/api/characterisation/analyse", files={"files": (DEMO.name, f.read(), "text/markdown")},
                          data={"meta": '{"product_name": "DemoVan 99"}', "domain": "feed"}).json()
    r = client.post("/api/characterisation/export.docx", json={"markdown": res["markdown"], "product_name": res["meta"]["product_name"]})
    assert r.status_code == 200
    d = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Characterisation of the additive — DemoVan 99" in text and "Data-gap analysis against the guidance" in text
    assert len(d.tables) >= 5
