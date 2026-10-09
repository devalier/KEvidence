"""Step 0 — substance/product characterisation of feed additives (chemical MVP).

Grounding model
---------------
Every check the module performs comes from a requirement catalogue
(data/guidance/feedap_2017_5023_chemical_requirements.json) derived from
EFSA FEEDAP (2017) "Guidance on the identity, characterisation and conditions
of use of feed additives", EFSA Journal 15(10):5023. Each requirement carries a
verification level. When the guidance text itself is loaded (PDF or text file in
the guidance directory), each requirement is anchored to the passage(s) of the
guidance that support it, and requirements that cannot be anchored are flagged.

Extraction is deterministic (regex/keyword) and every extracted value carries a
citation to the submitted document and page. Nothing here decides compliance:
statuses mean "evidence located" or "not located", for verification by the
scientific officer.

Confidentiality: documents are processed in memory for a single request and are
never written to disk. Optional LLM drafting is off unless explicitly enabled
per request, because it sends extracted excerpts to an external API.
"""

from __future__ import annotations

import datetime as _dt
import io
import ipaddress
import json
import logging
import os
import re
import socket
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Callable, Optional
from urllib.parse import urljoin, urlparse

GUIDANCE_DIR = Path(os.getenv("KEVIDENCE_GUIDANCE_DIR", str(Path(__file__).resolve().parent / "data" / "guidance")))
CATALOGUE_PATH = GUIDANCE_DIR / "feedap_2017_5023_chemical_requirements.json"
TEMPLATE_PATH = GUIDANCE_DIR / "characterisation_template.md"
# An institution can point this at its own (non-public) opinion template; placeholders are documented in
# data/guidance/characterisation_template.md. data/guidance/local/ is git-ignored for that purpose.
LOCAL_TEMPLATE_PATH = Path(os.getenv("KEVIDENCE_CHARACTERISATION_TEMPLATE", str(GUIDANCE_DIR / "local" / "characterisation_template.md")))
GUIDANCE_PDF_CANDIDATES = ("efsa_2017_5023.pdf", "j.efsa.2017.5023.pdf", "5023.pdf", "guidance_5023.pdf")

logging.getLogger("pypdf").setLevel(logging.ERROR)

MAX_FILE_BYTES = int(os.getenv("KEVIDENCE_MAX_FILE_BYTES", str(25 * 1024 * 1024)))
MAX_FILES = int(os.getenv("KEVIDENCE_MAX_FILES", "30"))
MAX_LINKED_DOCS = int(os.getenv("KEVIDENCE_MAX_LINKED_DOCS", "15"))
URL_TIMEOUT_S = 30.0
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".md", ".json", ".html", ".htm")

STATUS_LABELS = {
    "gap": "Not located (data gap)",
    "partial": "Partially addressed",
    "conditional": "Applicability to confirm",
    "manual_review": "Scientific-officer review",
    "evidence_located": "Evidence located",
    "not_applicable": "Not applicable",
}

NUMBER_WORDS = {
    "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7,
    "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12,
}


# ---------------------------------------------------------------------------
# Catalogue and guidance text
# ---------------------------------------------------------------------------

def load_catalogue(path: Path = CATALOGUE_PATH) -> dict[str, Any]:
    with open(path, encoding="utf-8") as f:
        catalogue = json.load(f)
    ids = [r["id"] for r in catalogue["requirements"]]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate requirement IDs in characterisation catalogue")
    tokens = {s["token"] for s in catalogue["template_sections"]}
    for req in catalogue["requirements"]:
        for key in ("guidance_section", "guidance_page", "quote", "guidance_anchors", "requirement"):
            if not req.get(key):
                raise ValueError(f"{req['id']} is missing '{key}' (every requirement must cite the guidance)")
        if req["template_section"] not in tokens:
            raise ValueError(f"{req['id']} references unknown template section {req['template_section']}")
        for pattern in req.get("patterns", []) + [p for s in req.get("sub_items", []) for p in s["patterns"]]:
            re.compile(pattern, re.I)
    return catalogue


def _normalise_ws(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


LIGATURES = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl", "–": "-", "—": "-", "’": "'", "‘": "'"}


def _fix_ligatures(text: str) -> str:
    for k, v in LIGATURES.items():
        text = text.replace(k, v)
    # The guidance PDF encodes "µ" in a symbol font that pypdf extracts as "l".
    return re.sub(r"(\d)\s?lm\b", "\\1 µm", text)


def _compact(text: str) -> tuple[str, list[int]]:
    """Lower-case, ligature-fixed text with all whitespace removed, plus a map
    from compact positions back to positions in the ligature-fixed text.
    PDF extraction inserts spurious spaces (e.g. 'speci ﬁcation'), so anchors
    are matched whitespace-insensitively."""
    chars, index = [], []
    for i, ch in enumerate(text):
        if not ch.isspace():
            chars.append(ch.lower())
            index.append(i)
    return "".join(chars), index


def load_guidance_pages(guidance_dir: Path = GUIDANCE_DIR) -> Optional[list[str]]:
    """Return the bundled guidance PDF as a list of page texts, or None if absent."""
    for name in GUIDANCE_PDF_CANDIDATES:
        pdf_path = guidance_dir / name
        if pdf_path.exists():
            doc = extract_document(pdf_path.name, pdf_path.read_bytes())
            if doc.pages and any(p.strip() for p in doc.pages):
                return doc.pages
    return None


def anchor_requirements(catalogue: dict[str, Any], pages: Optional[list[str]]) -> dict[str, dict[str, Any]]:
    """Locate each requirement's anchor phrases in the guidance text.

    Returns {requirement_id: {"status", "passages": [{"page", "anchor", "excerpt"}]}}.
    """
    result: dict[str, dict[str, Any]] = {}
    norm_pages = [_normalise_ws(_fix_ligatures(p)) for p in pages] if pages else None
    compact_pages = [_compact(p) for p in norm_pages] if norm_pages else None
    for req in catalogue["requirements"]:
        if not norm_pages:
            result[req["id"]] = {"status": "verified_primary", "passages": [],
                                 "note": f"Guidance text not loaded; catalogue quote from p.{req.get('guidance_page')} (section {req.get('guidance_section')})."}
            continue
        passages = []
        for anchor in req.get("guidance_anchors", []):
            needle, _ = _compact(_fix_ligatures(anchor))
            for page_no, (page, (compact, index)) in enumerate(zip(norm_pages, compact_pages), start=1):
                idx = compact.find(needle)
                if idx < 0:
                    continue
                start_raw = index[idx]
                end_raw = index[min(idx + len(needle), len(index)) - 1] + 1
                start = max(0, start_raw - 160)
                end = min(len(page), end_raw + 320)
                passages.append({
                    "page": page_no,
                    "anchor": anchor,
                    "excerpt": ("…" if start else "") + page[start:end] + ("…" if end < len(page) else ""),
                })
                break
        found_anchors = {p["anchor"] for p in passages}
        if passages and len(found_anchors) == len(req.get("guidance_anchors", [])):
            status = "anchored_in_guidance_text"
        else:
            status = "not_found_in_guidance_text"
        result[req["id"]] = {
            "status": status,
            "anchors_found": sorted(found_anchors),
            "anchors_missing": [a for a in req.get("guidance_anchors", []) if a not in found_anchors],
            "passages": passages[:3],
        }
    return result


def guidance_status(catalogue: Optional[dict[str, Any]] = None, guidance_dir: Path = GUIDANCE_DIR) -> dict[str, Any]:
    catalogue = catalogue or load_catalogue()
    pages = load_guidance_pages(guidance_dir)
    anchors = anchor_requirements(catalogue, pages)
    counts: dict[str, int] = {}
    for req in catalogue["requirements"]:
        key = anchors[req["id"]]["status"]
        counts[key] = counts.get(key, 0) + 1
    return {
        "guidance": catalogue["guidance"],
        "catalogue_id": catalogue["catalogue_id"],
        "catalogue_version": catalogue["catalogue_version"],
        "scope": catalogue["scope"],
        "guidance_text_loaded": bool(pages),
        "guidance_pages": len(pages) if pages else 0,
        "verification_levels": catalogue["verification_levels"],
        "verification_counts": counts,
        "template_sections": catalogue["template_sections"],
        "meta_options": catalogue.get("meta_options", {}),
        "related_guidance": catalogue.get("related_guidance", []),
        "requirements": [
            {
                "id": r["id"],
                "title": r["title"],
                "guidance_section": r["guidance_section"],
                "guidance_page": r["guidance_page"],
                "template_section": r["template_section"],
                "scope": r["scope"],
                "requirement": r["requirement"],
                "quote": r["quote"],
                "grounding": anchors[r["id"]],
            }
            for r in catalogue["requirements"]
        ],
    }


# ---------------------------------------------------------------------------
# Document ingestion
# ---------------------------------------------------------------------------

@dataclass
class Document:
    doc_id: str
    name: str
    source: str  # "upload" or URL
    pages: list[str] = field(default_factory=list)
    kind: str = "unknown"
    doc_type: str = "Unclassified"
    warnings: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return sum(len(p) for p in self.pages)


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.links: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self._skip += 1
        if tag == "a":
            for key, value in attrs:
                if key == "href" and value:
                    self.links.append(value)
        if tag in ("p", "br", "tr", "li", "div", "h1", "h2", "h3", "h4", "table"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript") and self._skip:
            self._skip -= 1
        if tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _extension(name: str) -> str:
    name = (name or "").lower().split("?")[0]
    return os.path.splitext(name)[1]


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = content.decode(enc)
            if enc == "utf-16" and "\x00" in text:
                continue
            return text
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _pdf_pages(content: bytes, warnings: list[str]) -> list[str]:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - dependency declared in requirements.txt
        warnings.append("pypdf is not installed; PDF text could not be extracted.")
        return []
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                warnings.append("PDF is encrypted and could not be opened without a password.")
                return []
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pages.append("")
                warnings.append("Text extraction failed on at least one PDF page.")
        if pages and sum(len(p.strip()) for p in pages) < 40 * len(pages):
            warnings.append("Very little extractable text: the PDF is probably scanned. OCR is required before this document can be assessed.")
        return pages
    except Exception as exc:
        warnings.append(f"PDF could not be parsed ({type(exc).__name__}).")
        return []


def _docx_pages(content: bytes, warnings: list[str]) -> list[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    except (zipfile.BadZipFile, KeyError):
        warnings.append("DOCX could not be opened.")
        return []
    xml = re.sub(r"</w:tc>", " | ", xml)
    xml = re.sub(r"</w:p>|<w:br/>|</w:tr>", "\n", xml)
    xml = re.sub(r"<w:tab/>", "\t", xml)
    text = re.sub(r"<[^>]+>", "", xml)
    text = (text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
            .replace("&quot;", '"').replace("&apos;", "'"))
    return [text]


def _xlsx_pages(content: bytes, warnings: list[str]) -> list[str]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        warnings.append(f"Spreadsheet could not be opened ({type(exc).__name__}).")
        return []
    pages = []
    for ws in wb.worksheets:
        lines = [f"[Sheet: {ws.title}]"]
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells))
        pages.append("\n".join(lines))
    wb.close()
    return pages


def html_to_text(content: bytes) -> tuple[str, list[str]]:
    parser = _HTMLText()
    parser.feed(_decode(content))
    text = re.sub(r"[ \t]+", " ", "".join(parser.parts))
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip(), parser.links


def extract_document(name: str, content: bytes, doc_id: str = "D?", source: str = "upload", content_type: str = "") -> Document:
    doc = Document(doc_id=doc_id, name=name or "unnamed", source=source)
    if len(content) > MAX_FILE_BYTES:
        doc.warnings.append(f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB limit and was not processed.")
        return doc
    ext = _extension(name)
    ctype = (content_type or "").lower()
    if ext == ".pdf" or "pdf" in ctype or content[:5] == b"%PDF-":
        doc.kind = "pdf"
        doc.pages = _pdf_pages(content, doc.warnings)
    elif ext == ".docx" or "wordprocessingml" in ctype:
        doc.kind = "docx"
        doc.pages = _docx_pages(content, doc.warnings)
    elif ext in (".xlsx", ".xlsm") or "spreadsheetml" in ctype:
        doc.kind = "xlsx"
        doc.pages = _xlsx_pages(content, doc.warnings)
    elif ext in (".html", ".htm") or "html" in ctype:
        doc.kind = "html"
        text, _ = html_to_text(content)
        doc.pages = [text]
    elif ext in (".csv", ".tsv", ".txt", ".md", ".json") or ctype.startswith("text/") or "json" in ctype:
        doc.kind = "text"
        doc.pages = [_decode(content)]
    else:
        doc.warnings.append(
            f"Unsupported file type '{ext or ctype or 'unknown'}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    doc.pages = [p.replace("\u00a0", " ") for p in doc.pages]
    doc.doc_type = classify_document(doc)
    return doc


DOC_TYPE_RULES = [
    ("Certificate of analysis", [r"certificat(e)?\s+(of|or|d')?\s*analy", r"\bCoA\b", r"analytical\s+certificate"]),
    ("Product specification", [r"product\s+specification", r"specification\s+sheet", r"technical\s+data\s+sheet"]),
    ("Safety data sheet", [r"safety\s+data\s+sheet", r"\bSDS\b", r"\bMSDS\b"]),
    ("Batch analysis report", [r"batch\s+(analysis|analyses|to\s*-?\s*batch)", r"batch-to-batch"]),
    ("Impurity / contaminant report", [r"heavy\s+metals", r"dioxins?", r"mycotoxins?", r"residual\s+solvents"]),
    ("Stability study", [r"stability\s+(study|studies|test)", r"shelf[\s-]*life"]),
    ("Homogeneity study", [r"homogeneity"]),
    ("Particle size / dusting study", [r"dusting\s+potential", r"particle\s+size", r"stauber"]),
    ("Manufacturing process description", [r"manufacturing\s+process", r"flow\s*(chart|diagram)", r"production\s+process"]),
    ("Technical dossier (Section II)", [r"section\s+II", r"identity,?\s+characterisation"]),
]


def classify_document(doc: Document) -> str:
    head = " ".join(doc.pages)[:20000]
    scores = []
    for label, patterns in DOC_TYPE_RULES:
        score = sum(len(re.findall(p, head, re.I)) for p in patterns)
        if score:
            scores.append((score, label))
    if not scores:
        return "Unclassified"
    scores.sort(reverse=True)
    return scores[0][1]


# --- URL retrieval (SSRF-guarded) -------------------------------------------

def _host_is_public(host: str) -> bool:
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        # DNS may only be resolvable by an outbound proxy; treat literal private
        # names as unsafe and let the HTTP client report failure otherwise.
        return not (host in ("localhost",) or host.endswith(".local") or host.endswith(".internal"))
    for info in infos:
        addr = ipaddress.ip_address(info[4][0])
        if (addr.is_private or addr.is_loopback or addr.is_link_local or addr.is_reserved
                or addr.is_multicast or addr.is_unspecified):
            return False
    return True


def validate_url(url: str) -> str:
    parsed = urlparse((url or "").strip())
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ValueError(f"Only http(s) URLs are accepted: {url!r}")
    try:
        ip = ipaddress.ip_address(parsed.hostname)
        if not ip.is_global:
            raise ValueError(f"URL host is not a public address: {url!r}")
    except ValueError as exc:
        if "not a public address" in str(exc):
            raise
        if not _host_is_public(parsed.hostname):
            raise ValueError(f"URL host resolves to a non-public address: {url!r}")
    return parsed.geturl()


def _fetch(url: str, client) -> tuple[bytes, str, str]:
    current = validate_url(url)
    for _ in range(4):
        with client.stream("GET", current, follow_redirects=False) as resp:
            if resp.status_code in (301, 302, 303, 307, 308) and resp.headers.get("location"):
                current = validate_url(urljoin(current, resp.headers["location"]))
                continue
            resp.raise_for_status()
            chunks, size = [], 0
            for chunk in resp.iter_bytes():
                size += len(chunk)
                if size > MAX_FILE_BYTES:
                    raise ValueError(f"Remote document exceeds {MAX_FILE_BYTES // (1024 * 1024)} MB: {current}")
                chunks.append(chunk)
            return b"".join(chunks), resp.headers.get("content-type", ""), current
    raise ValueError(f"Too many redirects: {url}")


def fetch_url_documents(url: str, start_index: int, client=None) -> tuple[list[Document], list[str]]:
    """Fetch a URL. HTML pages are kept as a document and linked dossier files
    (PDF/DOCX/XLSX) are followed, up to MAX_LINKED_DOCS."""
    import httpx

    notes: list[str] = []
    own_client = client is None
    client = client or httpx.Client(timeout=URL_TIMEOUT_S, headers={"User-Agent": "KEvidence/0.1 (characterisation)"})
    docs: list[Document] = []
    try:
        content, ctype, final_url = _fetch(url, client)
        name = os.path.basename(urlparse(final_url).path) or urlparse(final_url).hostname or "url"
        doc = extract_document(name, content, f"D{start_index}", final_url, ctype)
        docs.append(doc)
        if doc.kind == "html":
            _, links = html_to_text(content)
            doc_links = []
            for link in links:
                absolute = urljoin(final_url, link)
                if _extension(absolute) in (".pdf", ".docx", ".xlsx", ".xlsm") and absolute not in doc_links:
                    doc_links.append(absolute)
            if len(doc.pages[0] if doc.pages else "") < 500 and not doc_links:
                notes.append(
                    f"{final_url} returned little text and no document links. Dossier portals rendered by "
                    "JavaScript (e.g. Open EFSA) cannot be read by URL; download the documents and upload them."
                )
            for link in doc_links[:MAX_LINKED_DOCS]:
                try:
                    c2, t2, f2 = _fetch(link, client)
                    docs.append(extract_document(os.path.basename(urlparse(f2).path), c2,
                                                 f"D{start_index + len(docs)}", f2, t2))
                except Exception as exc:
                    notes.append(f"Linked document not retrieved ({link}): {exc}")
            if len(doc_links) > MAX_LINKED_DOCS:
                notes.append(f"{len(doc_links) - MAX_LINKED_DOCS} further linked documents were not followed (limit {MAX_LINKED_DOCS}).")
    except Exception as exc:
        notes.append(f"URL not retrieved ({url}): {exc}")
    finally:
        if own_client:
            client.close()
    return docs, notes


# ---------------------------------------------------------------------------
# Deterministic extraction
# ---------------------------------------------------------------------------

@dataclass
class Segment:
    doc_id: str
    page: int
    line: int
    text: str


def segment_documents(docs: list[Document]) -> list[Segment]:
    segments = []
    for doc in docs:
        for page_no, page in enumerate(doc.pages, start=1):
            for line_no, raw in enumerate(page.splitlines(), start=1):
                text = _normalise_ws(raw)
                if len(text) >= 2:
                    segments.append(Segment(doc.doc_id, page_no, line_no, text))
    return segments


def cite(seg: Segment) -> str:
    return f"{seg.doc_id} p.{seg.page}"


def cas_checksum_ok(cas: str) -> bool:
    m = re.fullmatch(r"(\d{2,7})-(\d{2})-(\d)", cas)
    if not m:
        return False
    digits = (m.group(1) + m.group(2))[::-1]
    return sum((i + 1) * int(d) for i, d in enumerate(digits)) % 10 == int(m.group(3))


def ec_checksum_ok(ec: str) -> bool:
    m = re.fullmatch(r"(\d{3})-(\d{3})-(\d)", ec)
    if not m:
        return False
    digits = m.group(1) + m.group(2)
    remainder = sum((i + 1) * int(d) for i, d in enumerate(digits)) % 11
    return remainder != 10 and remainder == int(m.group(3))


CAS_RE = re.compile(r"(?<![\d-])(\d{2,7}-\d{2}-\d)(?![\d-])")
EC_RE = re.compile(r"(?<![\d-])(\d{3}-\d{3}-\d)(?![\d-])")
SUBSCRIPTS = str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")
FORMULA_TOKEN_RE = re.compile(r"\b((?:[A-Z][a-z]?\d*){2,}(?:\s*[·•.]\s*\d*\s*(?:[A-Z][a-z]?\d*)+)?)\b")
ELEMENTS = set("H He Li Be B C N O F Ne Na Mg Al Si P S Cl Ar K Ca Sc Ti V Cr Mn Fe Co Ni Cu Zn Ga Ge As Se Br Kr Rb Sr Y Zr Nb Mo Tc Ru Rh Pd Ag Cd In Sn Sb Te I Xe Cs Ba La Ce Pr Nd Pm Sm Eu Gd Tb Dy Ho Er Tm Yb Lu Hf Ta W Re Os Ir Pt Au Hg Tl Pb Bi".split())
MW_RE = re.compile(r"(?:molecular|molar|formula)\s+(?:weight|mass)[^0-9]{0,40}?(\d{1,4}(?:[.,]\d{1,4})?)\s*(g\s*/\s*mol|g·mol|g mol|Da|daltons?)?", re.I)
SPEC_RE = re.compile(r"(≥|>=|≤|<=|not\s+less\s+than|not\s+more\s+than|min(?:imum|\.)?|max(?:imum|\.)?)\s*(\d+(?:[.,]\d+)?)\s*(%|g/kg|mg/kg)", re.I)
BATCH_ID_RE = re.compile(r"\b(?:batch|lot)(?![a-z])\s*(?:no\.?|number|nr\.?|n°|#|id|code)?\s*[:.]?\s*([A-Z0-9][A-Z0-9\-/.]*\d[A-Z0-9\-/.]*)", re.I)
BATCH_ROW_RE = re.compile(r"^(?:batch|lot)(?:es)?\s*(?:no\.?|number|nr\.?|n°|#|id|code)?\s*[:|]\s*(.+)$", re.I)
BATCH_COUNT_RE = re.compile(r"\b(\d{1,2}|one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve)\s+(?:independent\s+|different\s+|consecutive\s+|production\s+|representative\s+|commercial\s+)*(?:production\s+)?(?:batches|lots)\b", re.I)


def valid_formula(token: str) -> bool:
    token = token.translate(SUBSCRIPTS)
    parts = re.findall(r"([A-Z][a-z]?)(\d*)", token)
    if not parts:
        return False
    elements = [e for e, _ in parts]
    if not all(e in ELEMENTS for e in elements):
        return False
    # Require at least one stoichiometric number, to exclude acronyms such as "HPLC".
    return any(n for _, n in parts) and len(elements) >= 2


def find_cas_numbers(segments: list[Segment]) -> list[dict[str, Any]]:
    found: dict[str, dict[str, Any]] = {}
    for seg in segments:
        for m in CAS_RE.finditer(seg.text):
            cas = m.group(1)
            if EC_RE.fullmatch(cas):
                continue
            if not cas_checksum_ok(cas):
                if re.search(r"\bCAS\b", seg.text, re.I):
                    found.setdefault(f"invalid:{cas}", {"cas": cas, "valid": False, "citation": cite(seg), "context": seg.text})
                continue
            if cas not in found:
                found[cas] = {"cas": cas, "valid": True, "citation": cite(seg), "context": seg.text}
    return list(found.values())


def _value_after_label(text: str, label_re: str) -> Optional[str]:
    m = re.search(label_re + r"\s*(?:\([^)]*\))?\s*[:=–\-|]\s*(.+)", text, re.I)
    if m:
        value = m.group(1).strip(" |;")
        return value[:300] if value else None
    return None


def extract_identity(scope: list[Segment]) -> dict[str, Any]:
    """Pull identity fields from a list of segments; each value carries a citation."""
    out: dict[str, Any] = {}

    def put(key, value, seg):
        if value and key not in out:
            out[key] = {"value": value, "citation": cite(seg), "context": seg.text[:300]}

    for idx, seg in enumerate(scope):
        t = seg.text
        if re.search(r"IUPAC", t, re.I):
            put("iupac_name", _value_after_label(t, r"IUPAC(?:\s+name)?"), seg)
        if re.search(r"\bCAS\b", t, re.I):
            for m in CAS_RE.finditer(t):
                if cas_checksum_ok(m.group(1)):
                    put("cas", m.group(1), seg)
                    break
        if re.search(r"\bEC\b|EINECS", t):
            for m in EC_RE.finditer(t):
                if ec_checksum_ok(m.group(1)):
                    put("ec", m.group(1), seg)
                    break
        if re.search(r"(molecular|chemical|empirical)\s+formula", t, re.I):
            tail = re.split(r"formula", t, flags=re.I, maxsplit=1)[-1].translate(SUBSCRIPTS)
            for m in FORMULA_TOKEN_RE.finditer(tail):
                if valid_formula(m.group(1)):
                    put("molecular_formula", m.group(1).replace(" ", ""), seg)
                    break
        if re.search(r"structural\s+formula|chemical\s+structure|SMILES|InChI", t, re.I):
            value = _value_after_label(t, r"(?:structural\s+formula|chemical\s+structure|SMILES|InChI(?:Key)?)")
            put("structural_formula", value or "Structure referred to (figure/description) — verify in source", seg)
        mw = MW_RE.search(t)
        if mw:
            put("molecular_weight", f"{mw.group(1)} {mw.group(2) or 'g/mol (unit not stated)'}".strip(), seg)
        if re.search(r"synonym|other\s+names?|also\s+known\s+as", t, re.I):
            put("synonyms", _value_after_label(t, r"(?:synonyms?|other\s+names?|also\s+known\s+as)"), seg)
        if re.search(r"purity|\bassay\b|content\s+of", t, re.I):
            spec = SPEC_RE.search(t)
            if spec:
                put("purity", f"{spec.group(1)} {spec.group(2)} {spec.group(3)}", seg)
        if re.search(r"water\s+solubility|solubility\s+in\s+water|soluble\s+in\s+water", t, re.I):
            put("water_solubility", _value_after_label(t, r"(?:water\s+solubility|solubility\s+in\s+water)") or t[:200], seg)
        if re.search(r"log\s*K\s*ow|log\s*P\b|partition\s+coefficient", t, re.I):
            m = re.search(r"(?:log\s*K\s*ow|log\s*P|partition\s+coefficient)[^0-9\-−]{0,30}([\-−]?\d+(?:[.,]\d+)?)", t, re.I)
            put("log_kow", m.group(1).replace("−", "-") if m else t[:200], seg)
        if re.search(r"trade\s*name|proprietary\s+name|product\s+name", t, re.I):
            put("trade_name", _value_after_label(t, r"(?:trade\s*name|proprietary\s+name|product\s+name)"), seg)
        if re.search(r"generic\s+name|common\s+name", t, re.I):
            put("generic_name", _value_after_label(t, r"(?:generic|common)\s+name"), seg)
        if re.search(r"SMILES", t):
            m = re.search(r"SMILES[^:=|]*[:=|]\s*([A-Za-z0-9@+\-\[\]\(\)\\/=#$%.:]+)", t)
            put("smiles", m.group(1) if m else None, seg)
        if re.search(r"FLAVIS|FL[\s-]?no", t, re.I):
            m = re.search(r"\b(\d{2}\.\d{3})\b", t)
            put("flavis", m.group(1) if m else None, seg)
        if re.search(r"\bpKa\b|dissociation\s+constant", t, re.I):
            put("pka", _value_after_label(t, r"(?:pKa\w*|dissociation\s+constants?)") or t[:200], seg)
        if re.search(r"melting\s+point", t, re.I):
            put("melting_point", _value_after_label(t, r"melting\s+point") or t[:200], seg)
        if re.search(r"optical\s+rotation|specific\s+rotation", t, re.I):
            put("optical_rotation", _value_after_label(t, r"(?:specific\s+)?(?:optical\s+)?rotation") or t[:200], seg)
    return out


def detect_batches(segments: list[Segment]) -> dict[str, Any]:
    ids: dict[str, str] = {}
    stated: list[dict[str, Any]] = []
    for seg in segments:
        header = BATCH_ROW_RE.match(seg.text)
        if header:
            # Table row such as "Batch No. | 2301A | 2302B | 2303C"
            for cell in re.split(r"[|;,\t]|\s{2,}", header.group(1)):
                cell = cell.strip(" .")
                if re.fullmatch(r"[A-Z0-9][A-Z0-9\-/.]{0,24}", cell, re.I) and re.search(r"\d", cell):
                    ids.setdefault(cell.upper(), cite(seg))
        for m in BATCH_ID_RE.finditer(seg.text):
            batch_id = m.group(1).strip(".-/")
            if batch_id and batch_id.upper() not in ids and not re.fullmatch(r"(19|20)\d\d", batch_id):
                ids[batch_id.upper()] = cite(seg)
        for m in BATCH_COUNT_RE.finditer(seg.text):
            raw = m.group(1).lower()
            n = NUMBER_WORDS.get(raw, int(raw) if raw.isdigit() else 0)
            if n:
                stated.append({"count": n, "citation": cite(seg), "text": seg.text[:200]})
    max_stated = max((s["count"] for s in stated), default=0)
    return {
        "distinct_batch_ids": sorted(ids),
        "batch_id_citations": ids,
        "stated_counts": stated[:5],
        "best_estimate": max(len(ids), max_stated),
    }


# ---------------------------------------------------------------------------
# Requirement evaluation
# ---------------------------------------------------------------------------

HEADING_RE = re.compile(r"^(#+\s|\d+(\.\d+)*\.?\s+[A-Z][^:]{0,80}$)")
JUSTIFICATION_RE = re.compile(
    r"not\s+applicable|not\s+relevant|not\s+required|waiv|justif|omitted|omission|"
    r"no\s+(data|study|studies)\s+(is|are|was|were|has\s+been|have\s+been)\s+(provided|performed|conducted)",
    re.I,
)
YEAR_RE = re.compile(r"\b(19[89]\d|20[0-4]\d)\b")
MONTHS_RE = re.compile(r"\b(\d{1,2})\s*(months?|mo\.?)\b|\b(\d{1,3})\s*weeks?\b", re.I)
HOURS_RE = re.compile(r"\b(\d{1,3})\s*(h|hrs?|hours?)\b|\b(\d{1,2})\s*days?\b", re.I)
SUBSAMPLES_RE = re.compile(r"\b(\d{1,3}|ten|twelve|fifteen|twenty)\s+(?:sub-?\s*)?(?:samples|subsamples|replicates)\b", re.I)
WORD_NUM = {"ten": 10, "twelve": 12, "fifteen": 15, "twenty": 20}


def _compile(patterns: list[str]) -> list[re.Pattern]:
    return [re.compile(p, re.I) for p in patterns]


def _matches(patterns: list[re.Pattern], text: str) -> bool:
    return any(p.search(text) for p in patterns)


def _meta_list(meta: dict[str, Any], key: str) -> set[str]:
    value = meta.get(key) or []
    if isinstance(value, str):
        value = [value]
    return {str(v).strip().lower() for v in value if str(v).strip()}


def _applicable(applies_when: dict[str, Any], meta: dict[str, Any]) -> str:
    """Return 'yes', 'no', 'unknown' or 'manual' for a set of applicability conditions.

    Exemptions (`*_not`) apply only when the user has stated the triggering
    attribute; conditions that need a positive attribute (`*_any`, form, route)
    are 'unknown' until the attribute is entered.
    """
    if not applies_when:
        return "yes"
    form = (meta.get("product_form") or "").strip().lower()
    routes = _meta_list(meta, "routes")
    production = _meta_list(meta, "production_types")
    additive_type = _meta_list(meta, "additive_type")
    verdicts = []
    if "form" in applies_when:
        verdicts.append("unknown" if form in ("", "unknown") else ("yes" if form == applies_when["form"] else "no"))
    if "route" in applies_when:
        verdicts.append("unknown" if not routes else ("yes" if applies_when["route"] in routes else "no"))
    if "route_any" in applies_when:
        verdicts.append("unknown" if not routes else ("yes" if routes & set(applies_when["route_any"]) else "no"))
    if "production_type_any" in applies_when:
        verdicts.append("unknown" if not production else ("yes" if production & set(applies_when["production_type_any"]) else "no"))
    if "production_type_not" in applies_when:
        verdicts.append("no" if production and production <= set(applies_when["production_type_not"]) else "yes")
    if "additive_type_any" in applies_when:
        verdicts.append("unknown" if not additive_type else ("yes" if additive_type & set(applies_when["additive_type_any"]) else "no"))
    if "additive_type_not" in applies_when:
        verdicts.append("no" if additive_type & set(applies_when["additive_type_not"]) else "yes")
    if "components_min" in applies_when:
        verdicts.append("yes" if int(meta.get("_component_count") or 0) >= applies_when["components_min"] else "no")
    if "no" in verdicts:
        return "no"
    if "unknown" in verdicts:
        return "unknown"
    if applies_when.get("manual"):
        return "manual"
    return "yes"


def _page_neighbourhood(segments: list[Segment], hits: list[Segment], window: int = 12) -> list[Segment]:
    """Segments within `window` lines of any hit on the same page (tables of
    batch results usually sit just below the heading or label that matched)."""
    near: dict[tuple[str, int], list[int]] = {}
    for h in hits:
        near.setdefault((h.doc_id, h.page), []).append(h.line)
    return [s for s in segments
            if any(abs(s.line - line) <= window for line in near.get((s.doc_id, s.page), []))]


def _max_duration(segments: list[Segment], unit: str) -> Optional[tuple[float, str]]:
    best: Optional[tuple[float, str]] = None
    for seg in segments:
        regex = MONTHS_RE if unit == "months" else HOURS_RE
        for m in regex.finditer(seg.text):
            if unit == "months":
                value = float(m.group(1)) if m.group(1) else float(m.group(3)) / 4.345
            else:
                value = float(m.group(1)) if m.group(1) else float(m.group(3)) * 24
            if best is None or value > best[0]:
                best = (value, cite(seg))
    return best


def _check(label: str, met: Optional[bool], detail: str) -> dict[str, Any]:
    return {"label": label, "met": met, "detail": detail}


def evaluate_requirement(req: dict[str, Any], segments: list[Segment], meta: dict[str, Any]) -> dict[str, Any]:
    applicability = _applicable(req.get("applies_when", {}), meta)
    result: dict[str, Any] = {
        "id": req["id"],
        "title": req["title"],
        "guidance_section": req["guidance_section"],
        "guidance_page": req["guidance_page"],
        "section": req["template_section"],
        "requirement": req["requirement"],
        "quote": req["quote"],
        "applicability": applicability,
        "applicability_note": (req.get("applies_when") or {}).get("manual"),
        "evidence": [],
        "sub_items": [],
        "missing": [],
        "checks": [],
        "flags": [],
        "justification": [],
        "batch_check": None,
    }
    if req.get("conditional_note"):
        result["conditional_note"] = req["conditional_note"]
    if applicability == "no":
        result["status"] = "not_applicable"
        return result

    patterns = _compile(req.get("patterns", []))
    hits = [s for s in segments if _matches(patterns, s.text)]
    neighbourhood = _page_neighbourhood(segments, hits) if hits else []
    result["evidence"] = [{"citation": cite(s), "text": s.text[:300]} for s in hits[:8]]
    result["evidence_count"] = len(hits)

    if req.get("requires_meta") == "production_type" and not _meta_list(meta, "production_types"):
        result["missing"].append("Production type not entered — the minimum impurity set of §2.1.4 depends on it "
                                 "(chemical synthesis, fermentation, plant-derived, animal-derived, mineral).")

    production = _meta_list(meta, "production_types")
    required_missing = []
    for sub in req.get("sub_items", []):
        sub_app = _applicable(sub.get("applies_when", {}), meta)
        if sub_app == "no":
            continue
        sub_patterns = _compile(sub["patterns"])
        sub_hits = [s for s in hits if _matches(sub_patterns, s.text)] or [s for s in segments if _matches(sub_patterns, s.text)]
        sub_hits.sort(key=lambda seg: bool(HEADING_RE.match(seg.text)))  # content lines before headings
        located = bool(sub_hits)
        notes = [note for ptype, note in (sub.get("note_when") or {}).items() if ptype in production]
        entry = {
            "label": sub["label"],
            "required": bool(sub.get("required")) and sub_app in ("yes", "manual"),
            "applicability": sub_app,
            "located": located,
            "citations": sorted({cite(s) for s in sub_hits})[:6],
            "snippet": f"{sub_hits[0].text[:160]}" if sub_hits else "",
            "note": "; ".join(notes) or None,
            "field": sub.get("field"),
        }
        result["sub_items"].append(entry)
        if not located:
            suffix = f" ({entry['note']})" if entry["note"] else ""
            if entry["required"]:
                required_missing.append(sub["label"])
                result["missing"].append(f"{sub['label']} — not located in the submitted documents{suffix}")
            elif sub_app == "unknown" and sub.get("required"):
                result["missing"].append(f"{sub['label']} — not located; required if applicable (enter product form / route / production type)")
            else:
                result["missing"].append(f"{sub['label']} — not located (provide where relevant){suffix}")

    failed_check = False
    if hits:
        if req.get("min_batches"):
            batches = detect_batches(neighbourhood)
            ok = batches["best_estimate"] >= req["min_batches"]
            result["batch_check"] = {
                "required_min": req["min_batches"],
                "detected": batches["best_estimate"],
                "batch_ids": batches["distinct_batch_ids"][:20],
                "stated_counts": batches["stated_counts"],
                "met": ok,
            }
            if not ok:
                failed_check = True
                result["missing"].append(
                    f"At least {req['min_batches']} batches required (§{req['guidance_section']}); "
                    f"{batches['best_estimate']} detected on the pages carrying this evidence"
                )
        if req.get("recency_years"):
            years = [(int(m.group(1)), cite(s)) for s in neighbourhood for m in YEAR_RE.finditer(s.text)]
            limit = _dt.date.today().year - req["recency_years"]
            if not years:
                result["checks"].append(_check("Analyses produced within the last 5 years", None,
                                               "No dates located near the analytical data — confirm the date of the certificates of analysis."))
            else:
                latest = max(years)
                ok = latest[0] >= limit
                result["checks"].append(_check(
                    "Analyses produced within the last 5 years", ok,
                    f"Most recent year detected: {latest[0]} ({latest[1]})" + ("" if ok else f" — older than {limit}")))
                if not ok:
                    failed_check = True
                    result["missing"].append(f"Analytical data appear older than 5 years (latest year detected {latest[0]})")
        for key, unit, label in (("min_months", "months", "Study duration"), ("min_hours", "hours", "Study duration")):
            if req.get(key):
                found = _max_duration(neighbourhood, unit)
                if not found:
                    result["checks"].append(_check(f"{label} ≥ {req[key]} {unit}", None, "Duration not located — verify in the study report."))
                else:
                    ok = found[0] >= req[key]
                    result["checks"].append(_check(f"{label} ≥ {req[key]} {unit}", ok, f"Longest duration detected: {found[0]:g} {unit} ({found[1]})"))
                    if not ok:
                        failed_check = True
                        result["missing"].append(f"Duration below the guidance minimum of {req[key]} {unit} (§{req['guidance_section']})")
        if req.get("min_subsamples"):
            counts = []
            for s in neighbourhood:
                for m in SUBSAMPLES_RE.finditer(s.text):
                    raw = m.group(1).lower()
                    counts.append((WORD_NUM.get(raw, int(raw) if raw.isdigit() else 0), cite(s)))
            if not counts:
                result["checks"].append(_check(f"≥ {req['min_subsamples']} subsamples", None, "Number of subsamples not located."))
            else:
                best = max(counts)
                ok = best[0] >= req["min_subsamples"]
                result["checks"].append(_check(f"≥ {req['min_subsamples']} subsamples", ok, f"{best[0]} detected ({best[1]})"))
                if not ok:
                    failed_check = True
                    result["missing"].append(f"Fewer than {req['min_subsamples']} subsamples detected")
        if req.get("red_flag_patterns"):
            flagged = [s for s in neighbourhood if _matches(_compile(req["red_flag_patterns"]), s.text)]
            if flagged:
                result["flags"].append({"message": req["red_flag_message"],
                                        "citations": sorted({cite(s) for s in flagged})[:6]})
        if req.get("pilot_batch_patterns"):
            pilot = [s for s in neighbourhood if _matches(_compile(req["pilot_batch_patterns"]), s.text)]
            if pilot:
                result["flags"].append({"message": "Pilot batches referred to — acceptable only where commercial batches are not yet "
                                                   "available and they represent the intended manufacturing process (§2.1.3).",
                                        "citations": sorted({cite(s) for s in pilot})[:6]})

    if req.get("manual_review"):
        result["status"] = "manual_review"
    elif not hits:
        result["status"] = "gap"
        if not any("not located" in m for m in result["missing"]):
            result["missing"].append("No evidence located in the submitted documents")
    elif required_missing or failed_check:
        result["status"] = "partial"
    else:
        result["status"] = "evidence_located"
    if applicability in ("unknown", "manual") and result["status"] in ("gap", "partial"):
        result["status"] = "conditional"
    if req.get("requires_meta") == "production_type" and not production and result["status"] != "gap":
        result["status"] = "conditional"

    if result["status"] in ("gap", "partial", "conditional"):
        just = [s for s in segments if JUSTIFICATION_RE.search(s.text) and _matches(patterns, s.text)]
        result["justification"] = [{"citation": cite(s), "text": s.text[:240]} for s in just[:3]]
    return result


def _component_scope(segments: list[Segment], component: dict[str, Any], window: int = 25) -> list[Segment]:
    needles = [n.lower() for n in (component.get("name"), component.get("cas")) if n]
    if not needles:
        return []
    scope_keys: set[tuple[str, int, int]] = set()
    for i, seg in enumerate(segments):
        low = seg.text.lower()
        if any(n in low for n in needles):
            for j in range(i, min(len(segments), i + window)):
                if segments[j].doc_id == seg.doc_id and segments[j].page == seg.page:
                    scope_keys.add((segments[j].doc_id, segments[j].page, segments[j].line))
    return [s for s in segments if (s.doc_id, s.page, s.line) in scope_keys]


def infer_components(segments: list[Segment]) -> list[dict[str, Any]]:
    comps = []
    for entry in find_cas_numbers(segments):
        if not entry["valid"]:
            continue
        context = entry["context"]
        label = re.split(r"\(?\s*CAS", context, flags=re.I)[0].strip(" :|-–(") or ""
        label = label if 2 <= len(label) <= 80 else ""
        comps.append({"name": label, "cas": entry["cas"], "role": "Detected (role to confirm)", "auto": True,
                      "detected_from": entry["citation"]})
    return comps[:8]


def _default_components(segments: list[Segment], meta: dict[str, Any]) -> list[dict[str, Any]]:
    components = infer_components(segments)
    product = meta.get("product_name")
    if "plant_derived" in _meta_list(meta, "production_types") and product and \
            not any(product.lower() in (c.get("name") or "").lower() for c in components):
        # A plant extract is characterised as a whole, by its constituents/marker compounds (§2.1.3, §2.2.1.1).
        source = (meta.get("_inferred") or {}).get("product_name", {}).get("citation")
        components.insert(0, {"name": product, "cas": "", "role": "Active substance (preparation of plant origin)",
                              "auto": True, "detected_from": source or "entered"})
    if not components and product:
        source = (meta.get("_inferred") or {}).get("product_name", {}).get("citation")
        components = [{"name": product, "cas": "", "role": "Additive (identity to confirm)", "auto": True,
                       "detected_from": source or "entered"}]
    return components


def evaluate_components(catalogue: dict[str, Any], segments: list[Segment], meta: dict[str, Any]) -> list[dict[str, Any]]:
    components = [c for c in meta.get("components", []) if (c.get("name") or c.get("cas"))]
    if not components:
        components = _default_components(segments, meta)
    component_reqs = [r for r in catalogue["requirements"] if r["scope"] == "component"]
    out = []
    single = len(components) <= 1
    for comp in components:
        scope = segments if single else _component_scope(segments, comp)
        identity = extract_identity(scope)
        if not comp.get("name"):
            for key in ("generic_name", "iupac_name"):
                if identity.get(key):
                    comp = {**comp, "name": str(identity[key]["value"])[:80]}
                    break
        reqs = []
        for req in component_reqs:
            evaluation = evaluate_requirement(req, scope, meta)
            for sub in evaluation["sub_items"]:
                if sub.get("field") and sub["field"] in identity:
                    sub["value"] = identity[sub["field"]]
            reqs.append(evaluation)
        cas_given = (comp.get("cas") or "").strip()
        cas_note = None
        if cas_given:
            cas_note = "CAS check digit valid" if cas_checksum_ok(cas_given) else "CAS check digit INVALID — verify the number"
            if identity.get("cas") and identity["cas"]["value"] != cas_given:
                cas_note += f"; documents state {identity['cas']['value']} ({identity['cas']['citation']})"
        out.append({
            "component": comp,
            "scope_segments": len(scope),
            "identity": identity,
            "cas_check": cas_note,
            "requirements": reqs,
        })
    return out


def _fermentation_flags(segments: list[Segment]) -> list[dict[str, str]]:
    compiled = _compile([r"ferment", r"production\s+(strain|organism)", r"corynebacterium", r"bacillus\s+\w+", r"aspergillus",
                         r"saccharomyces", r"escherichia\s+coli\s+\w*\d", r"\b(DSM|KCCM|CGMCC|NRRL|ATCC|NCIMB)\s*\d"])
    return [{"citation": cite(s), "text": s.text[:200]} for s in segments if _matches(compiled, s.text)][:5]


# ---------------------------------------------------------------------------
# Identification from the documents (product, form, production type, markers)
# ---------------------------------------------------------------------------
# Scientific officers upload documents and expect KEvidence to identify what
# they describe. Anything entered in the form takes precedence; everything
# inferred here is labelled "detected — confirm" with its citation.

COA_UNIT = r"(%|g/kg|mg/kg|µg/kg|μg/kg|ppm|ppb|CFU/g|cfu/g|mg/g|g/100\s*g)"
COA_LINE_RE = re.compile(
    r"^(?P<label>[^≥≤<>]*?)\s*(?P<cmp>≥|≤|>=|<=|NLT|NMT|min\.?|max\.?)\s*(?P<spec>\d+(?:[.,]\d+)?)\s*(?P<unit>" + COA_UNIT + r")"
    r"\s+(?P<lt>[<≤])?\s*(?P<res>\d+(?:[.,]\d+)?)\s*(?P<runit>" + COA_UNIT + r")?\s*$", re.I)
PRODUCT_NAME_RE = re.compile(r"\b([A-Z][\w-]*(?:\s+[a-z][\w-]*){0,3}\s+(?:extract|essential\s+oil|tincture|oleoresin|concentrate))\b")
PRODUCT_LABEL_RE = re.compile(r"(?:product\s+name|trade\s*name|name\s+of\s+the\s+(?:additive|product)|designation)\s*[:|]\s*(.{3,80})", re.I)
FORM_RULES = [("solid", re.compile(r"\b(powder|granul\w*|crystal\w*|microgranul\w*|pellet\w*|flakes?|beads?)\b", re.I)),
              ("liquid", re.compile(r"\b(liquid|solution|syrup|emulsion|suspension)\b", re.I))]
PRODUCTION_RULES = [
    ("plant_derived", re.compile(r"\bextract\b|botanical|essential\s+oil|tincture|oleoresin|\b[A-Z][a-z]+aceae\b|plant\s+origin|"
                                 r"\b(leaves|pomace|grape|herb\w*)\b", re.I)),
    ("fermentation", re.compile(r"ferment|production\s+(strain|organism)|\b(DSM|KCCM|CGMCC|NRRL|ATCC|NCIMB)\s*\d", re.I)),
    ("chemical_synthesis", re.compile(r"chemical(ly)?\s+synthes|synthesi[sz]ed\b|by\s+synthesis", re.I)),
    ("mineral", re.compile(r"trace\s+elements?|\bchelate\b|\b(zinc|copper|iron|manganese|selenium|cobalt|iodine)\b.{0,30}"
                           r"\b(oxide|sulphate|sulfate|carbonate|chloride|chelate|hydroxychloride)\b", re.I)),
]
AUTH_REG_RE = re.compile(r"(?:Regulation\s*\((?:EU|EC)\)\s*(?:No\.?\s*)?|\b(?:EU|EC)\s+(?:No\.?\s*)?)(\d{2,4}/\d{1,4})", re.I)
MARKER_EXCLUDE = re.compile(r"water|moisture|loss\s+on\s+drying|\bash\b|lead|cadmium|mercury|arsenic|dioxin|salmonella|"
                            r"enterobacter|yeast|mould|mold|coli|aerobic|density|particle|solvent|methanol|ethanol", re.I)


def _resolve_coa_label(segments: list[Segment], i: int, inline: str) -> str:
    """CoA tables often put the label on the preceding line(s); bilingual CoAs give the
    translation last, so the immediately preceding lines are used."""
    label = inline.strip(" :|")
    if len(re.sub(r"[^A-Za-z]", "", label)) >= 2:
        return label
    seg = segments[i]
    parts: list[str] = []
    j = i - 1
    while j >= 0 and segments[j].doc_id == seg.doc_id and segments[j].page == seg.page and len(parts) < 3:
        text = segments[j].text.strip()
        if COA_LINE_RE.match(text) or not re.search(r"[A-Za-z]", text):
            break
        parts.insert(0, text)
        if not text[:1].islower():
            break
        j -= 1
    return " ".join(parts)


def extract_coa_lines(segments: list[Segment]) -> list[dict[str, Any]]:
    """'Total polyphenols ≥ 80 % 81.4 %' → label, specification and result (one batch)."""
    out = []
    for i, seg in enumerate(segments):
        m = COA_LINE_RE.match(seg.text)
        if not m:
            continue
        label = _resolve_coa_label(segments, i, m.group("label"))
        if not label:
            continue
        cmp_raw = m.group("cmp").lower().rstrip(".")
        comparator = {"≥": "≥", ">=": "≥", "nlt": "≥", "min": "≥", "≤": "≤", "<=": "≤", "nmt": "≤", "max": "≤"}[cmp_raw]
        out.append({"label": label, "comparator": comparator, "spec": float(m.group("spec").replace(",", ".")),
                    "spec_raw": m.group("spec"), "unit": re.sub(r"\s+", " ", m.group("unit")),
                    "result": float(m.group("res").replace(",", ".")), "result_raw": m.group("res"), "lt": bool(m.group("lt")),
                    "result_unit": re.sub(r"\s+", " ", m.group("runit") or m.group("unit")), "segment": seg})
    return out


def _first_match(segments: list[Segment], regex: re.Pattern, group: int = 0) -> Optional[tuple[str, Segment]]:
    for seg in segments:
        m = regex.search(seg.text)
        if m:
            return m.group(group), seg
    return None


def infer_product(docs: list[Document], segments: list[Segment]) -> dict[str, Any]:
    inferred: dict[str, Any] = {}
    named = _first_match(segments, PRODUCT_LABEL_RE, 1)
    if not named:
        heads = [s for s in segments if s.line <= 6 and s.page == 1]
        named = _first_match(heads, PRODUCT_NAME_RE, 1) or _first_match(segments, PRODUCT_NAME_RE, 1)
    if named:
        inferred["product_name"] = {"value": named[0].strip(" .:|"), "citation": cite(named[1])}
    appearance = [s for s in segments if re.search(r"appearance|aspect|physical\s+(form|state)|form\b", s.text, re.I)]
    for form, regex in FORM_RULES:
        hit = _first_match(appearance + segments, regex)
        if hit:
            inferred["product_form"] = {"value": form, "citation": cite(hit[1]), "evidence": hit[1].text[:120]}
            break
    types = []
    for ptype, regex in PRODUCTION_RULES:
        hit = _first_match(segments, regex)
        if hit:
            types.append({"value": ptype, "citation": cite(hit[1]), "evidence": hit[1].text[:120]})
    if types:
        inferred["production_types"] = types
    reg = next(((m.group(1), s) for s in segments if re.search(r"authori[sz]|autori[sz]|zulassung", s.text, re.I)
                for m in [AUTH_REG_RE.search(s.text)] if m), None)
    if reg:
        inferred["authorisation_reference"] = {"value": f"Regulation (EU) {reg[0]}" if "/" in reg[0] else reg[0],
                                               "citation": cite(reg[1]), "evidence": reg[1].text[:160]}
    markers = []
    for row in extract_coa_lines(segments):
        if row["comparator"] == "≥" and not MARKER_EXCLUDE.search(row["label"]) and \
                row["label"].lower() not in {m["label"].lower() for m in markers}:
            markers.append({"label": row["label"], "specification": f"≥ {row['spec_raw']} {row['unit']}",
                            "citation": cite(row["segment"])})
    if markers:
        inferred["markers"] = markers
    batches = detect_batches(segments)
    if batches["distinct_batch_ids"]:
        inferred["batches"] = {"ids": batches["distinct_batch_ids"], "citations": batches["batch_id_citations"]}
    return inferred


def apply_inference(meta: dict[str, Any], inferred: dict[str, Any]) -> dict[str, Any]:
    """Fill blanks in the user's input with what the documents show; record what was inferred."""
    meta = dict(meta)
    used: dict[str, Any] = {}
    if not meta.get("product_name") and inferred.get("product_name"):
        meta["product_name"] = inferred["product_name"]["value"]
        used["product_name"] = inferred["product_name"]
    if not meta.get("product_form") and inferred.get("product_form"):
        meta["product_form"] = inferred["product_form"]["value"]
        used["product_form"] = inferred["product_form"]
    if not _meta_list(meta, "production_types") and inferred.get("production_types"):
        meta["production_types"] = sorted({t["value"] for t in inferred["production_types"]})
        used["production_types"] = inferred["production_types"]
    if not meta.get("authorisation") and inferred.get("authorisation_reference"):
        meta["authorisation"] = f"Applicant refers to an authorising regulation: {inferred['authorisation_reference']['value']}"
        used["authorisation"] = inferred["authorisation_reference"]
    meta["_inferred"] = used
    return meta


def analyse_dossier(docs: list[Document], meta: dict[str, Any], catalogue: Optional[dict[str, Any]] = None,
                    guidance_pages: Optional[list[str]] = None, fetch_notes: Optional[list[str]] = None) -> dict[str, Any]:
    catalogue = catalogue or load_catalogue()
    segments = segment_documents(docs)
    meta = dict(meta or {})
    meta["routes"] = sorted(_meta_list(meta, "routes"))
    meta["production_types"] = sorted(_meta_list(meta, "production_types"))
    inferred = infer_product(docs, segments)
    meta = apply_inference(meta, inferred)
    meta["production_types"] = sorted(_meta_list(meta, "production_types"))
    entered = [c for c in meta.get("components", []) if (c.get("name") or c.get("cas"))]
    meta["_component_count"] = len(entered) or len(_default_components(segments, meta))
    additive_reqs = [r for r in catalogue["requirements"] if r["scope"] == "additive"]
    additive_results = [evaluate_requirement(r, segments, meta) for r in additive_reqs]
    components = evaluate_components(catalogue, segments, meta)
    anchors = anchor_requirements(catalogue, guidance_pages)
    for res in additive_results:
        res["grounding"] = anchors[res["id"]]
    for comp in components:
        for res in comp["requirements"]:
            res["grounding"] = anchors[res["id"]]

    all_results = additive_results + [r for c in components for r in c["requirements"]]
    summary: dict[str, int] = {k: 0 for k in STATUS_LABELS}
    for res in all_results:
        summary[res["status"]] = summary.get(res["status"], 0) + 1

    scope_flags = []
    fermentation = _fermentation_flags(segments)
    if fermentation and "fermentation" not in meta["production_types"]:
        scope_flags.append({
            "flag": "Production by fermentation indicated in the documents",
            "detail": "Select 'fermentation' as production type so that the §2.1.4 fermentation-specific impurity requirements "
                      "(spent growth medium, absence of production organism and antimicrobial activity, LPS, DNA) are applied. "
                      "Characterisation of the production strain itself follows the FEEDAP microorganism guidance "
                      "(EFSA Journal 2018;16(3):5206) and is outside this MVP.",
            "evidence": fermentation,
        })
    elif "fermentation" in meta["production_types"]:
        scope_flags.append({
            "flag": "Fermentation product",
            "detail": "Production-strain characterisation (§2.2.1.2, §2.2.2.2) follows the FEEDAP microorganism guidance "
                      "(EFSA Journal 2018;16(3):5206) and is outside this chemical-characterisation MVP.",
            "evidence": fermentation,
        })
    if not components:
        scope_flags.append({
            "flag": "Substance not identified from the documents",
            "detail": "KEvidence found no product name, CAS number or constituent specification in the documents. Check that the "
                      "documents contain text (scanned PDFs need OCR), or enter the product name and components under "
                      "'Optional details' and run again (§2.1.3, §2.2.1.1).",
            "evidence": [],
        })
    unreadable = [d for d in docs if not d.char_count]
    if unreadable:
        scope_flags.append({
            "flag": "Documents without extractable text",
            "detail": "Absence of evidence for these documents may reflect extraction limits rather than true data gaps: "
                      + ", ".join(f"{d.doc_id} {d.name}" for d in unreadable),
            "evidence": [],
        })

    meta["formulations"] = [str(f).strip() for f in (meta.get("formulations") or []) if str(f).strip()]
    component_names = [c["component"].get("name") or "" for c in components]
    opinion = build_opinion_tables(docs, segments, meta, component_names)
    meta_out = {k: v for k, v in meta.items() if not k.startswith("_")}
    detected = {"inferred_and_used": meta.get("_inferred", {}), "all": inferred}
    result = {
        "detected": detected,
        "generated_at": _dt.datetime.now(_dt.timezone.utc).isoformat(timespec="seconds"),
        "meta": meta_out,
        "catalogue": {"id": catalogue["catalogue_id"], "version": catalogue["catalogue_version"],
                      "guidance": catalogue["guidance"]},
        "guidance_text_loaded": bool(guidance_pages),
        "documents": [
            {"doc_id": d.doc_id, "name": d.name, "source": d.source, "kind": d.kind, "doc_type": d.doc_type,
             "pages": len(d.pages), "characters": d.char_count, "warnings": d.warnings}
            for d in docs
        ],
        "fetch_notes": fetch_notes or [],
        "cas_numbers_found": find_cas_numbers(segments),
        "components": components,
        "requirements": additive_results,
        "summary": summary,
        "scope_flags": scope_flags,
        "opinion": opinion,
    }
    result["markdown"] = render_markdown(result, catalogue)
    return result


# ---------------------------------------------------------------------------
# Opinion-template tables (FEEDAP characterisation section layout)
# ---------------------------------------------------------------------------
# The FEEDAP opinion characterisation section reports, per product/active
# substance: specifications; batch-to-batch variation as average (range) with
# the number of batches in []; substance-related impurities as ranges; and, in
# Appendix A, other impurities and physico-chemical/technological properties.
# "<" marks values below the LOQ and "-" parameters not analysed. The values
# below are extracted as stated in the dossier, each with its source citation.

UNIT_PATTERN = (r"(%\s*(?:w/w|DM)?|g/kg(?:\s*DM)?|mg/kg(?:\s*DM)?|µg/kg|μg/kg|ug/kg|ng\s*(?:WHO-?)?TEQ/kg|ng/kg|"
                r"CFU/g|cfu/g|log10\s*CFU/g|mg/m\s*(?:3|³)|g/cm\s*(?:3|³)|g/mL|g/ml|g/L|g/l|µm|μm|ppm|ppb)")
SOLVENTS = ["dichloromethane", "methanol", "ethanol", "acetone", "hexane", "n-hexane", "heptane", "toluene", "isopropanol",
            "2-propanol", "ethyl acetate", "acetonitrile", "chloroform", "tetrahydrofuran", "cyclohexane", "xylene",
            "methyl tert-butyl ether", "dimethylformamide", "pyridine", "benzene"]
OTHER_IMPURITY_TERMS = ["lead", "cadmium", "mercury", "arsenic", "fluorine", "fluoride", "nickel", "chromium", "copper",
                        "dioxins and dioxin-like PCBs", "dioxins", "dioxin", "dl-PCBs", "PCDD/F", "aflatoxin B1", "aflatoxins",
                        "ochratoxin A", "zearalenone", "deoxynivalenol", "fumonisins", "T-2", "mycotoxins", "pesticides",
                        "Salmonella", "Enterobacteriaceae", "Escherichia coli", "E. coli", "yeasts and moulds",
                        "yeasts and filamentous fungi", "filamentous fungi", "moulds", "yeasts", "Bacillus cereus",
                        "total aerobic count", "total viable count", "lipopolysaccharides", "endotoxins"]
PHYS_TERMS = ["dusting potential", "bulk density", "tapped density", "density", "particle size", "particles below 10 µm",
              "particles below 100 µm", "particles below 50 µm", "particles below 1 µm", "D10", "D50", "D90", "viscosity",
              "vapour pressure", "pH", "solubility", "melting point", "specific weight"]
SUBSTANCE_RELATED_CUES = re.compile(r"residual|synthes|impurit|related\s+substance|by-?product|intermediate|solvent", re.I)
ACTIVE_CUES = re.compile(r"\bassay\b|purity|content|active\s+substance|identified|total\s+amount|polyphenol|proanthocyan|"
                         r"anthocyan|flavon|flavan|tannin|glucoside|catechin|resveratrol|carotenoid|terpen|saponin|alkaloid|"
                         r"thymol|carvacrol|curcumin|marker", re.I)
OTHER_PARAM_CUES = re.compile(r"loss\s+on\s+drying|moisture|water\s+content|^water\b|\bash\b|sulphated\s+ash|sulfated\s+ash|"
                              r"specific\s+(optical\s+)?rotation|\bpH\b|chloride|sulphate|sulfate", re.I)
SPEC_LINE_CUES = re.compile(r"specification|specified|composition|assay|purity|content|\bspec\b", re.I)
METHOD_CUES = re.compile(r"\b(HPLC(?:-[A-Z]+)?|LC-MS(?:/MS)?|GC(?:-[A-Z]+)?|ICP-(?:MS|OES|AES)|AAS|titration|"
                         r"ion[- ]exchange chromatography|Karl Fischer|IC|NMR|IR)\b")
CELL_SPLIT_RE = re.compile(r"\s*\|\s*|\t+|\s{2,}")
VALUE_CELL_RE = re.compile(r"^(<|≤|<=)?\s*(\d+(?:[.,]\d+)?)\s*(?:%|mg/kg|g/kg)?$")
ND_CELL_RE = re.compile(r"^(n\.?d\.?|not\s+detected|absent|negative|<\s*LO[DQ]|below\s+LO[DQ]|<LOQ|<LOD)$", re.I)
NA_CELL_RE = re.compile(r"^(-|–|n\.?a\.?|not\s+analy[sz]ed)$", re.I)


def _parse_cell(cell: str) -> Optional[dict[str, Any]]:
    cell = cell.strip()
    m = VALUE_CELL_RE.match(cell)
    if m:
        return {"raw": cell, "value": float(m.group(2).replace(",", ".")), "lt": bool(m.group(1)), "na": False}
    if ND_CELL_RE.match(cell):
        return {"raw": cell, "value": None, "lt": True, "na": False}
    if NA_CELL_RE.match(cell):
        return {"raw": cell, "value": None, "lt": False, "na": True}
    return None


def _unit_from_label(label: str) -> str:
    units = re.findall(UNIT_PATTERN, label)
    return re.sub(r"\s+", " ", units[-1]).strip() if units else ""


def _clean_label(label: str) -> str:
    label = re.sub(r"\(\s*" + UNIT_PATTERN + r"\s*\)", "", label)
    label = re.sub(r",\s*" + UNIT_PATTERN + r"\s*\)", ")", label)
    return re.sub(r"\s+", " ", label).strip(" :;,-")


def _categorise(label: str, line: str, component_names: list[str]) -> str:
    low = f"{label} {line}".lower()
    if any(s in low for s in SOLVENTS) or SUBSTANCE_RELATED_CUES.search(label):
        return "substance_related"
    if any(t.lower() in label.lower() for t in OTHER_IMPURITY_TERMS):
        return "other_impurities"
    if any(t.lower() in label.lower() for t in PHYS_TERMS if t != "pH") or re.search(r"\bpH\b", label):
        return "physchem" if not OTHER_PARAM_CUES.search(label) or "density" in label.lower() else "batch_other"
    if OTHER_PARAM_CUES.search(label):
        return "batch_other"
    if ACTIVE_CUES.search(label) or any(n and n.lower() in label.lower() for n in component_names):
        return "batch_active"
    if SUBSTANCE_RELATED_CUES.search(line):
        return "substance_related"
    return "other"


def _row(label: str, unit: str, cells: list[dict[str, Any]], seg: "Segment", category: str, n_stated: Optional[int] = None,
         source: str = "table") -> dict[str, Any]:
    numeric = [c["value"] for c in cells if c["value"] is not None and not c["lt"]]
    lt = [c for c in cells if c["lt"]]
    analysed = [c for c in cells if not c["na"]]
    return {
        "label": label, "unit": unit, "category": category, "values": [c["raw"] for c in cells],
        "n": n_stated or (len(analysed) if source in ("table", "coa") else 0), "n_stated": bool(n_stated), "numeric": numeric, "below_loq": len(lt),
        "lt_values": [c["value"] for c in lt if c["value"] is not None], "not_analysed": len(cells) - len(analysed),
        "citation": cite(seg), "doc_id": seg.doc_id, "source": source,
    }


def _term_regex(term: str) -> re.Pattern:
    number = r"(?:[<≤]\s*)?\d+(?:[.,]\d+)?"
    return re.compile(
        rf"\b{re.escape(term)}\b(?:\s*\([^)]*\))?[^0-9<≤|;]{{0,40}}?({number}(?:\s*(?:,|and|–|-|to)\s*{number})*)\s*{UNIT_PATTERN}",
        re.I if term not in ("pH", "D10", "D50", "D90") else 0,
    )


NAMED_VALUE_RE = re.compile(r"(?P<name>[A-Za-z][A-Za-z0-9 ,'\-()]{2,60}?)\s*(?::|=)?\s*(?P<lt><|≤)?\s*(?P<value>\d+(?:[.,]\d+)?)\s*"
                            + r"(?P<unit>" + UNIT_PATTERN + r")")
NAMED_ND_RE = re.compile(r"(?P<name>[A-Za-z][A-Za-z0-9 '\-]{2,60}?)\s+(?:was\s+)?not\s+detected\s*\((?:LOD|LOQ)\s*(?P<lt>)(?P<value>\d+(?:[.,]\d+)?)\s*"
                         + r"(?P<unit>" + UNIT_PATTERN + r")\)", re.I)
TERM_PATTERNS = [(t, _term_regex(t)) for t in sorted(set(SOLVENTS + OTHER_IMPURITY_TERMS + PHYS_TERMS), key=len, reverse=True)]


def extract_parameter_rows(segments: list[Segment], component_names: list[str]) -> list[dict[str, Any]]:
    """Batch tables ("label | v1 | v2 | v3") and analyte statements ("lead < 0.5 mg/kg in three batches")."""
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    coa_segments = set()
    for coa in extract_coa_lines(segments):
        seg = coa["segment"]
        coa_segments.add((seg.doc_id, seg.page, seg.line))
        key = (seg.doc_id, coa["label"].lower())
        if key in seen:
            continue
        seen.add(key)
        cell = {"raw": f"{'<' if coa['lt'] else ''}{coa['result_raw']}", "value": coa["result"], "lt": coa["lt"], "na": False}
        rows.append(_row(coa["label"], coa["result_unit"], [cell], seg, _categorise(coa["label"], seg.text, component_names),
                         source="coa"))
    for seg in segments:
        if (seg.doc_id, seg.page, seg.line) in coa_segments:
            continue
        cells = [c for c in CELL_SPLIT_RE.split(seg.text) if c.strip()]
        if len(cells) >= 3 and re.search(r"[A-Za-z]", cells[0]) and not BATCH_ID_RE.match(cells[0]) \
                and not re.match(r"^(batch|lot|parameter|analyte|sample)s?\b", cells[0], re.I):
            parsed = [_parse_cell(c) for c in cells[1:]]
            if all(parsed) and sum(1 for p in parsed if not p["na"]) >= 2:
                label = _clean_label(cells[0])
                key = (seg.doc_id, label.lower())
                if label and key not in seen:
                    seen.add(key)
                    rows.append(_row(label, _unit_from_label(cells[0]), parsed, seg,
                                     _categorise(cells[0], seg.text, component_names)))
                continue
        count = BATCH_COUNT_RE.search(seg.text)
        n_stated = None
        if count:
            raw = count.group(1).lower()
            n_stated = NUMBER_WORDS.get(raw, int(raw) if raw.isdigit() else None)
        taken: list[tuple[int, int]] = []
        if SUBSTANCE_RELATED_CUES.search(seg.text):
            body = seg.text.split(":", 1)[1] if ":" in seg.text else seg.text
            for clause in re.split(r"[;,]", body):
                m = NAMED_ND_RE.search(clause) or NAMED_VALUE_RE.search(clause)
                if not m or any(t.lower() in m.group("name").lower() for t in OTHER_IMPURITY_TERMS + PHYS_TERMS):
                    continue
                label = _clean_label(m.group("name"))
                key = (seg.doc_id, label.lower())
                if not label or key in seen or len(label) > 60:
                    continue
                seen.add(key)
                below = bool(m.group("lt")) or "not detected" in m.group(0).lower()
                cell = {"raw": f"{'<' if below else ''}{m.group('value')}", "value": float(m.group("value").replace(",", ".")),
                        "lt": below, "na": False}
                rows.append(_row(label, re.sub(r"\s+", " ", m.group("unit")).strip(), [cell], seg, "substance_related",
                                 n_stated, source="statement"))
        for term, regex in TERM_PATTERNS:
            for m in regex.finditer(seg.text):
                if any(m.start() < e and s < m.end() for s, e in taken):
                    continue
                taken.append((m.start(), m.end()))
                parts = re.split(r"\s*(?:,|and|–|-|to)\s*(?=[<≤]?\s*\d)", m.group(1))
                parsed = [p for p in (_parse_cell(x) for x in parts) if p]
                is_range = bool(re.search(r"\d\s*(?:–|-|to)\s*[<≤]?\s*\d", m.group(1)))
                label = term if term[0].isupper() or term in ("pH",) else term.capitalize()
                key = (seg.doc_id, label.lower())
                if not parsed or key in seen:
                    continue
                seen.add(key)
                row = _row(label, re.sub(r"\s+", " ", m.group(2)).strip(), parsed, seg,
                           _categorise(term, seg.text, component_names), n_stated, source="statement")
                if is_range:
                    row["stated_range"] = True
                    row["n"] = n_stated or 0
                rows.append(row)
    return rows


def _fmt(v: float) -> str:
    return f"{v:.4g}" if abs(v) < 1e5 else f"{v:.3e}"


def format_row_value(row: dict[str, Any], average: bool) -> str:
    """Average (range) [n] for batch-to-batch variation; range [n] otherwise. '<' = below LOQ; '-' = not analysed."""
    n = f" [{row['n']}]" if row["n"] else " [n not stated]"
    numeric, lts = row["numeric"], row["lt_values"]
    if not numeric and not row["below_loq"]:
        return "-"
    if not numeric:
        return (f"<{_fmt(max(lts))}" if lts else "<LOQ") + n
    lo, hi = min(numeric), max(numeric)
    low_txt = f"<{_fmt(max(lts))}" if row["below_loq"] and lts else ("<LOQ" if row["below_loq"] else _fmt(lo))
    if average and not row["below_loq"] and len(numeric) > 1 and not row.get("stated_range"):
        decimals = max((len(v.split(".")[1]) if "." in v else 0) for v in (str(x).replace(",", ".") for x in row["values"]))
        return f"{sum(numeric) / len(numeric):.{decimals}f} ({_fmt(lo)}–{_fmt(hi)}){n}"
    if len(numeric) == 1 and not row["below_loq"]:
        return f"{row['values'][0].strip()}{n}"
    return f"{low_txt}–{_fmt(hi)}{n}" if (row["below_loq"] or lo != hi) else f"{_fmt(lo)}{n}"


def extract_specifications(segments: list[Segment]) -> list[dict[str, Any]]:
    specs, seen = [], set()
    coa_seen: set[tuple[str, str, str, str]] = set()
    for row in extract_coa_lines(segments):
        key = (row["label"].lower(), row["comparator"], row["spec_raw"], row["unit"])
        if key in coa_seen:
            continue
        coa_seen.add(key)
        specs.append({"label": row["label"], "comparator": row["comparator"], "raw_value": row["spec_raw"], "also_stated": [],
                      "value": row["spec"], "unit": row["unit"], "text": f"{row['comparator']} {row['spec_raw']} {row['unit']}",
                      "method": None, "citation": cite(row["segment"]), "source": "coa"})
    for seg in segments:
        if not SPEC_LINE_CUES.search(seg.text):
            continue
        for clause in re.split(r";", seg.text):
            for m in SPEC_RE.finditer(clause):
                before = clause[:m.start()]
                label = _clean_label(re.split(r":", before)[-1]) or _clean_label(before)
                label = re.sub(r"^(specification|composition|spec)\s*", "", label, flags=re.I).strip(" :") or "Active substance"
                comparator = {"not less than": "≥", "min": "≥", "min.": "≥", "minimum": "≥", ">=": "≥",
                              "not more than": "≤", "max": "≤", "max.": "≤", "maximum": "≤", "<=": "≤"}.get(m.group(1).lower(), m.group(1))
                key = (comparator, m.group(2), m.group(3))
                method = METHOD_CUES.search(clause)
                if key in seen:
                    existing = next(sp for sp in specs if (sp["comparator"], sp["raw_value"], sp["unit"]) == key)
                    existing["also_stated"].append(f"{label} ({cite(seg)})")
                    if method and not existing["method"]:
                        existing["method"] = method.group(1)
                    continue
                seen.add(key)
                specs.append({"label": label, "comparator": comparator, "raw_value": m.group(2), "also_stated": [],
                              "value": float(m.group(2).replace(",", ".")),
                              "unit": m.group(3), "text": f"{comparator} {m.group(2)} {m.group(3)}",
                              "method": method.group(1) if method else None, "citation": cite(seg)})
    return specs


def _spec_compliance(rows: list[dict[str, Any]], specs: list[dict[str, Any]]) -> list[str]:
    facts = []
    generic = {"total", "content", "sum", "amount", "level", "other", "free"}
    for spec in specs:
        labels = [spec["label"]] + [a.rsplit(" (", 1)[0] for a in spec.get("also_stated", [])]
        exact = [r for r in rows if r["label"].lower() in {lab.lower() for lab in labels}]
        words = {w for lab in labels for w in re.findall(r"[a-z]{4,}", lab.lower())} - generic
        for row in rows:
            if row["category"] not in ("batch_active", "batch_other") or not row["numeric"]:
                continue
            if exact:
                if row not in exact:
                    continue
            elif not (words & (set(re.findall(r"[a-z]{4,}", row["label"].lower())) - generic)):
                continue
            if spec["comparator"] == "≥":
                ok = sum(1 for v in row["numeric"] if v >= spec["value"])
            elif spec["comparator"] == "≤":
                ok = sum(1 for v in row["numeric"] if v <= spec["value"]) + row["below_loq"]
            else:
                continue
            total = len(row["numeric"]) + (row["below_loq"] if spec["comparator"] == "≤" else 0)
            facts.append(f"{row['label']}: {ok}/{total} batch values meet the specification {spec['text']} "
                         f"({row['citation']}; specification {spec['citation']})")
    return list(dict.fromkeys(facts))


def _doc_formulation_map(docs: list[Document], formulations: list[str]) -> dict[str, Optional[str]]:
    mapping: dict[str, Optional[str]] = {}
    for doc in docs:
        text = " ".join(doc.pages).lower()
        counts = [(text.count(f.lower()), f) for f in formulations if f]
        counts = [c for c in counts if c[0]]
        mapping[doc.doc_id] = max(counts)[1] if counts else None
    return mapping


def _first_line(segments: list[Segment], pattern: str) -> Optional[Segment]:
    regex = re.compile(pattern, re.I)
    return next((s for s in segments if regex.search(s.text)), None)


def fermentation_fields(segments: list[Segment]) -> dict[str, dict[str, Any]]:
    """Fields the opinion template asks for on viable cells and DNA of the production strain."""
    out: dict[str, dict[str, Any]] = {}
    for key, cue in (("viable_cells", r"viable\s+cells|production\s+(strain|organism).{0,40}(absence|detected)"),
                     ("dna", r"\bDNA\b")):
        lines = [s for s in segments if re.search(cue, s.text, re.I)]
        if not lines:
            continue
        text = " ".join(s.text for s in lines)
        batches = BATCH_COUNT_RE.search(text)
        replicate = re.search(r"\b(duplicate|triplicate|quadruplicate)\b", text, re.I)
        sample = re.search(r"(\d+(?:[.,]\d+)?)\s*(g|gram|grams|mL)\s+(?:per\s+sample|of\s+(?:the\s+)?(?:product|sample))", text, re.I)
        fields = {
            "batches": batches.group(0) if batches else None,
            "replicates": replicate.group(1) if replicate else None,
            "sample_size": f"{sample.group(1)} {sample.group(2)} per sample" if sample else None,
            "result": (re.search(r"(no\s+(viable\s+cells|DNA)[^.]*|not\s+detected[^.]*|absen[ct][^.]*)", text, re.I) or [None])[0],
            "citations": sorted({cite(s) for s in lines})[:5],
        }
        if key == "dna":
            primers = re.search(r"(\d+)\s*bp", text)
            lod = re.search(r"(?:limit\s+of\s+detection|LOD)[^0-9]{0,30}(\d+(?:[.,]\d+)?)\s*(ng|pg|µg|μg)", text, re.I)
            fields["amplicon"] = f"{primers.group(1)} bp" if primers else None
            fields["lod"] = f"{lod.group(1)} {lod.group(2)} per gram" if lod else None
            fields["method"] = "PCR" if re.search(r"\bPCR\b", text) else None
        out[key] = fields
    return out


def build_opinion_tables(docs: list[Document], segments: list[Segment], meta: dict[str, Any],
                         component_names: list[str]) -> dict[str, Any]:
    formulations = [f for f in meta.get("formulations", []) if f]
    doc_map = _doc_formulation_map(docs, formulations) if formulations else {}
    rows = extract_parameter_rows(segments, component_names)
    specs = extract_specifications(segments)
    columns: dict[str, dict[str, Any]] = {}

    def merge(col_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """One certificate of analysis per batch: combine same parameter across documents into one row."""
        merged: dict[tuple[str, str, str], dict[str, Any]] = {}
        out = []
        for r in col_rows:
            key = (r["label"].lower(), r["unit"], r["category"])
            if r["source"] not in ("coa", "table") or key not in merged:
                merged.setdefault(key, r)
                out.append(r)
                continue
            m = merged[key]
            if r["doc_id"] in m.get("doc_ids", {m["doc_id"]}):
                continue
            m.update({"values": m["values"] + r["values"], "numeric": m["numeric"] + r["numeric"],
                      "below_loq": m["below_loq"] + r["below_loq"], "lt_values": m["lt_values"] + r["lt_values"],
                      "not_analysed": m["not_analysed"] + r["not_analysed"], "n": m["n"] + r["n"],
                      "doc_ids": m.get("doc_ids", {m["doc_id"]}) | {r["doc_id"]},
                      "citation": f"{m['citation']}; {r['citation']}"})
        return out

    def column(name: str, doc_ids: Optional[set[str]]) -> dict[str, Any]:
        segs = [s for s in segments if doc_ids is None or s.doc_id in doc_ids]
        col_rows = merge([dict(r) for r in rows if doc_ids is None or r["doc_id"] in doc_ids])
        col_specs = [sp for sp in specs if doc_ids is None or sp["citation"].split()[0] in doc_ids]
        form = _first_line(segs, r"physical\s+(form|state)|appearance|\b(powder|granul\w*|liquid|crystalline|microgranul\w*)\b")
        return {"name": name, "rows": col_rows, "specs": col_specs,
                "authorisation_reference": infer_product(docs, segs).get("authorisation_reference") if segs else None,
                "method_statement": (lambda m: {"text": m.text[:160], "citation": cite(m)} if m else None)(
                    _first_line(segs, r"(results?\s+obtained|determined|analy[sz]ed)\s+.{0,40}(analytical\s+)?methods?|"
                                      r"analytical\s+methods?\s+(described|as\s+described|according)")), "compliance": _spec_compliance(col_rows, col_specs),
                "physical_form": {"text": form.text[:160], "citation": cite(form)} if form else None,
                "documents": sorted({r["doc_id"] for r in col_rows} | {sp["citation"].split()[0] for sp in col_specs})}

    if formulations:
        unattributed = {d.doc_id for d in docs if not doc_map.get(d.doc_id)}
        table1 = column("Active substance / product", unattributed) if unattributed else None
        for f in formulations:
            columns[f] = column(f, {d for d, name in doc_map.items() if name == f})
    else:
        table1 = column(meta.get("product_name") or "Additive", None)
    return {"table1": table1, "table2": columns, "doc_formulation_map": doc_map, "all_rows": rows, "specifications": specs,
            "fermentation": fermentation_fields(segments)}


# ---------------------------------------------------------------------------
# Opinion-template rendering (§1.1.1, Appendix A, §1.1.2)
# ---------------------------------------------------------------------------

def _doc_names(result: dict[str, Any], doc_ids: set[str]) -> str:
    names = [f"{d['doc_id']} {d['name']}" for d in result["documents"] if d["doc_id"] in doc_ids]
    return "; ".join(names) or "[not located — Annex file names to be added]"


def _table_rows_md(col: dict[str, Any], categories: tuple[str, ...], average: bool) -> list[tuple[str, str, str]]:
    out = []
    for r in col["rows"]:
        if r["category"] in categories:
            label = f"{r['label']} ({r['unit']})" if r["unit"] and r["unit"] not in r["label"] else r["label"]
            out.append((label, format_row_value(r, average), r["citation"]))
    return out


def _table1_md(col: dict[str, Any], title: str) -> list[str]:
    lines = [f"**Table 1:** Data on specifications, batch-to-batch variation and substance-related impurities of {title}. "
             "The data presented are average values and (range) for batch-to-batch variation and ranges for all other "
             "parameters. The number of batches analysed per parameter or group of parameters is indicated in [].", "",
             "| Parameter | Value | Source |", "|---|---|---|", "| **Specifications**⁽¹⁾ | | |"]
    specs = col["specs"]
    lines += [f"| {_md_escape(sp['label'])}{(' (also stated as: ' + _md_escape(', '.join(sp['also_stated'])) + ')') if sp.get('also_stated') else ''}"
              f" | {_md_escape(sp['text'])} | {sp['citation']} |" for sp in specs] or \
             ["| Active substance (units) | **[not located]** | — |"]
    lines.append("| **Batch-to-batch variation** | | |")
    b2b = _table_rows_md(col, ("batch_active", "batch_other"), True)
    lines += [f"| {_md_escape(a)} | {_md_escape(b)} | {c} |" for a, b, c in b2b] or ["| Active substance (units) | **[not located]** | — |"]
    lines.append("| **Substance-related impurities** | | |")
    imp = _table_rows_md(col, ("substance_related",), False)
    lines += [f"| {_md_escape(a)} | {_md_escape(b)} | {c} |" for a, b, c in imp] or ["| Residual solvents / synthesis impurities (units) | **[not located]** | — |"]
    methods = sorted({sp["method"] for sp in specs if sp.get("method")})
    if not methods and col.get("method_statement"):
        methods = [f"\"{col['method_statement']['text']}\" ({col['method_statement']['citation']})"]
    reg = col.get("authorisation_reference")
    lines += ["", "Abbreviations: DM: dry matter.  ",
              f"⁽¹⁾ Method of analysis: {', '.join(methods) if methods else '[not located]'} // specifications set in the authorising "
              f"regulation: {(reg['value'] + ' (as stated in ' + reg['citation'] + ' — confirm)') if reg else '[to be completed, where applicable]'}.  ",
              "<: below the limit of quantification; -: not analysed.", ""]
    return lines


def _table2_md(result: dict[str, Any], cols: dict[str, dict[str, Any]], title: str) -> list[str]:
    names = list(cols)
    head = "| | " + " | ".join(_md_escape(n) for n in names) + " |"
    sep = "|---|" + "---|" * len(names)
    lines = [f"**Table 2:** Data on the specifications, batch-to-batch variation and substance-related impurities of preparations "
             f"containing {title}. The data presented are average values and (range) for batch-to-batch variation and ranges "
             "for all other parameters. The number of batches analysed per parameter or group of parameters is indicated in [].",
             "", head, sep]

    def section(label: str, categories: tuple[str, ...], average: bool, spec: bool = False):
        lines.append(f"| **{label}** |" + " |" * len(names))
        labels: list[str] = []
        per_col: dict[str, dict[str, str]] = {}
        for n in names:
            per_col[n] = {}
            if spec:
                for sp in cols[n]["specs"]:
                    per_col[n].setdefault(sp["label"], sp["text"])
            else:
                for a, b, _c in _table_rows_md(cols[n], categories, average):
                    per_col[n].setdefault(a, b)
            for k in per_col[n]:
                if k not in labels:
                    labels.append(k)
        if not labels:
            lines.append("| **[not located]** |" + " |" * len(names))
        for k in labels:
            lines.append(f"| {_md_escape(k)} | " + " | ".join(_md_escape(per_col[n].get(k, "-")) for n in names) + " |")

    section("Specifications⁽¹⁾", (), False, spec=True)
    lines.append("| **Physical form** | " + " | ".join(
        _md_escape(cols[n]["physical_form"]["text"]) if cols[n]["physical_form"] else "**[not located]**" for n in names) + " |")
    section("Batch-to-batch variation", ("batch_active", "batch_other"), True)
    section("Substance-related impurities", ("substance_related",), False)
    lines += ["", "Abbreviations: DM: dry matter.  ", "⁽¹⁾ Method of analysis: [to be completed].  ",
              "<: below the limit of quantification; -: not analysed.", ""]
    unattributed = [d for d, f in result["opinion"]["doc_formulation_map"].items() if not f]
    if unattributed:
        lines += [f"*Documents not attributed to a preparation (name not found in the document): {', '.join(unattributed)}.*", ""]
    return lines


def _detected_md(result: dict[str, Any]) -> str:
    det = result.get("detected", {})
    found, used = det.get("all", {}), det.get("inferred_and_used", {})
    rows = ["| Element | Identified from the documents | Source | Used in this analysis |", "|---|---|---|---|"]

    def add(label, item, key):
        if not item:
            rows.append(f"| {label} | **[not identified]** | — | — |")
            return
        value = item["value"] if isinstance(item, dict) else item
        rows.append(f"| {label} | {_md_escape(value)} | {item.get('citation', '—') if isinstance(item, dict) else '—'} | "
                    f"{'yes (detected — confirm)' if key in used else 'no — value entered by the user takes precedence'} |")
    add("Product / additive", found.get("product_name"), "product_name")
    add("Physical form", found.get("product_form"), "product_form")
    types = found.get("production_types") or []
    if types:
        for t in types:
            rows.append(f"| Production type | {t['value'].replace('_', ' ')} (\"{_md_escape(t['evidence'][:80])}\") | {t['citation']} | "
                        f"{'yes (detected — confirm)' if 'production_types' in used else 'no — entered value used'} |")
    else:
        rows.append("| Production type | **[not identified]** | — | — |")
    add("Authorising regulation referred to", found.get("authorisation_reference"), "authorisation")
    batches = found.get("batches")
    rows.append(f"| Batches | {', '.join(batches['ids'][:12]) if batches else '**[not identified]**'} | "
                f"{', '.join(sorted(set(batches['citations'].values()))[:5]) if batches else '—'} | yes |")
    for m in found.get("markers", []):
        rows.append(f"| Constituent / marker with specification | {_md_escape(m['label'])} {_md_escape(m['specification'])} | {m['citation']} | yes |")
    comps = ", ".join(f"{c['component'].get('name') or c['component'].get('cas')} ({c['component'].get('role')})" for c in result["components"])
    rows.append(f"| Component(s) characterised | {_md_escape(comps) or '**[none]**'} | — | yes |")
    return "\n".join(["Everything below was read from the uploaded documents. Values entered under *Optional details* override "
                      "detected ones; confirm or correct them there and run again.", ""] + rows)


def _opinion_111_md(result: dict[str, Any]) -> str:
    op = result["opinion"]
    meta = result["meta"]
    active = next((c for c in result["components"] if re.search(r"active", c["component"].get("role") or "", re.I)),
                  result["components"][0] if result["components"] else None)
    active_name = (active["component"].get("name") or active["component"].get("cas")) if active else "[active substance]"
    t1, t2 = op["table1"], op["table2"]
    all_cols = ([t1] if t1 else []) + list(t2.values())
    specs = [sp for col in all_cols for sp in col["specs"]]
    spec_txt = "; ".join(f"{sp['label']} {sp['text']}" for sp in specs) or "**[specification not located]**"
    rows = [r for col in all_cols for r in col["rows"]]
    def docs_of(categories):
        return {d for r in rows if r["category"] in categories for d in r.get("doc_ids", {r["doc_id"]})}
    b2b_docs = docs_of(("batch_active", "batch_other"))
    imp_docs = docs_of(("substance_related",))
    app_docs = docs_of(("other_impurities", "physchem", "other"))
    tables = "Table 1" + (" and Table 2" if t2 and t1 else "") if t1 else ("Table 2" if t2 else "Table 1")
    lines = [f"The specifications of the feed additive are: {_md_escape(spec_txt)}.", "",
             (f"Current authorisation: {_md_escape(meta['authorisation'])}"
              + (f" (detected in {result['detected']['inferred_and_used']['authorisation']['citation']} — confirm)"
                 if result.get("detected", {}).get("inferred_and_used", {}).get("authorisation") else " (as entered)")
              + ". **[If authorised: state the authorised minimum content.]**" if meta.get("authorisation") else
              "The additive is currently authorised with a minimum content of **[to be completed, or delete if this is a new application]**."),
             "",
             f"The data provided by the applicant on the batch-to-batch variation[^b2b] and substance-related impurities[^imp] "
             f"of the additive are reported in {tables}. Data on other impurities and physico-chemical and technological "
             f"properties[^phys] are reported in Appendix A.", ""]
    if t1:
        lines += _table1_md(t1, f"the product/active substance {active_name}")
    if t2:
        lines += _table2_md(result, t2, active_name)
    facts = [f for col in all_cols for f in col["compliance"]]
    flags = [f["message"] for r in result["requirements"] for f in r.get("flags", [])]
    lines += ["**Facts located for the scientific officer's conclusion** (not a conclusion):", ""]
    lines += [f"- {_md_escape(f)}" for f in facts] or ["- No batch values could be compared automatically with a specification."]
    solvent_rows = [r for r in rows if r["category"] == "substance_related" and any(s in r["label"].lower() for s in SOLVENTS)]
    if solvent_rows:
        lines.append("- Residual solvents reported: " + "; ".join(f"{r['label']} ({r['unit']}) {format_row_value(r, False)} ({r['citation']})"
                                                                  for r in solvent_rows)
                     + " — compare with the VICH GL18 limits cited in guidance §2.1.4.")
    else:
        lines.append("- No residual-solvent results located (guidance §2.1.4 requires residual solvents to be identified and quantified).")
    lines += [f"- ⚠ {_md_escape(f)}" for f in dict.fromkeys(flags)]
    gaps = [r for r in result["requirements"] if r["guidance_section"].startswith("2.1") and r["status"] in ("gap", "partial")]
    lines += [f"- Open point §{r['guidance_section']} {r['title']}: {_md_escape('; '.join(r['missing'][:2]))}" for r in gaps]
    lines += ["", "> **[Scientific officer to conclude]** on compliance with the specifications (set by the applicant or in the "
              "authorising regulation), on whether the microbial contamination and impurity levels raise concern, and on residual "
              "solvents relative to the VICH limits — or to identify levels that warrant monitoring during manufacturing.", ""]
    if "fermentation" in meta.get("production_types", []):
        ferm = op["fermentation"]
        lines += ["**Production strain in the final product (fermentation products):**", ""]
        vc, dna = ferm.get("viable_cells"), ferm.get("dna")

        def field(d, k):
            return _md_escape(d.get(k)) if d and d.get(k) else "**[not located]**"
        lines += [f"- Viable cells of the production strain: batches {field(vc, 'batches')}; replicates {field(vc, 'replicates')}; "
                  f"sample size {field(vc, 'sample_size')}; result {field(vc, 'result')}"
                  + (f" ({', '.join(vc['citations'])})" if vc else "") + ".",
                  f"- DNA of the production strain: method {field(dna, 'method')}; batches {field(dna, 'batches')}; replicates "
                  f"{field(dna, 'replicates')}; sample size {field(dna, 'sample_size')}; amplicon {field(dna, 'amplicon')}; "
                  f"LOD {field(dna, 'lod')}; result {field(dna, 'result')}" + (f" ({', '.join(dna['citations'])})" if dna else "") + ".",
                  "", "*Guidance §2.1.4 requires the absence of production organisms to be confirmed and, for GMM or strains carrying "
                  "AMR genes, the absence of their DNA to be demonstrated. The methodology details (replicates, sample size, LOD) follow "
                  "the EFSA Scientific Committee requirements referenced by the opinion template, not the 2017 guidance.*", ""]
    lines += [f"[^b2b]: Batch-to-batch variation: {_doc_names(result, b2b_docs)}",
              f"[^imp]: Substance-related impurities / residual solvents: {_doc_names(result, imp_docs)}",
              f"[^phys]: Other impurities and physical properties: {_doc_names(result, app_docs)}"]
    return "\n".join(lines)


def _appendix_a_md(result: dict[str, Any]) -> str:
    op = result["opinion"]
    cols = ([op["table1"]] if op["table1"] else []) + list(op["table2"].values())
    lines = ["**Table A.1:** Other impurities and physico-chemical and technological properties of the additive "
             "(ranges; number of batches in []). <: below the limit of quantification; -: not analysed.", "",
             "| Group | Parameter | Value | Product | Source |", "|---|---|---|---|---|"]
    group_names = {"other_impurities": "Other impurities", "physchem": "Physico-chemical / technological", "other": "Other parameters"}
    count = 0
    for col in cols:
        for r in col["rows"]:
            if r["category"] in group_names:
                count += 1
                label = f"{r['label']} ({r['unit']})" if r["unit"] and r["unit"] not in r["label"] else r["label"]
                lines.append(f"| {group_names[r['category']]} | {_md_escape(label)} | {_md_escape(format_row_value(r, False))} | "
                             f"{_md_escape(col['name'])} | {r['citation']} |")
    if not count:
        lines.append("| — | **[no values located]** | | | |")
    lines += ["", "**Technological properties (stability, homogeneity) — evidence located:**", ""]
    tech = [r for r in result["requirements"] if r["guidance_section"].startswith("2.4") and r["status"] != "not_applicable"]
    for r in tech:
        subs = [sb for sb in r.get("sub_items", []) if sb["located"] and sb.get("snippet")]
        ev = "; ".join(f"[{sb['citations'][0]}] {sb['snippet'][:140]}" for sb in subs[:2]) or "**[not located]**"
        lines.append(f"- §{r['guidance_section']} {r['title']} — {_status_md(r['status'])}: {_md_escape(ev)}")
    return "\n".join(lines)


def _opinion_112_md(result: dict[str, Any]) -> str:
    if "fermentation" not in result["meta"].get("production_types", []):
        return "*Not applicable unless the additive is, or is produced by, a microorganism.*"
    return ("**[Outside this chemical-characterisation MVP — to be completed by the scientific officer.]** Characterisation of the "
            "production strain (deposition, taxonomic identification by WGS, genetic modification, antimicrobial resistance genes and "
            "susceptibility, toxins and virulence factors, antimicrobial activity) follows the FEEDAP microorganism guidance "
            "(EFSA Journal 2018;16(3):5206) referred to in guidance §2.2.1.2 and §2.2.2.2.")


# ---------------------------------------------------------------------------
# Markdown rendering
# ---------------------------------------------------------------------------

STATUS_ICONS = {"evidence_located": "✅", "partial": "🟠", "gap": "❌", "not_applicable": "➖",
                "conditional": "❔", "manual_review": "👁"}


def _md_escape(text: Any) -> str:
    return str(text if text is not None else "").replace("|", "\\|").replace("\n", " ").strip()


def _status_md(status: str) -> str:
    return f"{STATUS_ICONS.get(status, '')} {STATUS_LABELS.get(status, status)}".strip()


def _grounding_md(res: dict[str, Any]) -> str:
    g = res.get("grounding") or {}
    status = g.get("status", "verified_primary")
    if status == "anchored_in_guidance_text":
        pages = ", ".join(sorted({f"p.{p['page']}" for p in g.get("passages", [])}))
        return f"anchored in guidance text ({pages})"
    if status == "not_found_in_guidance_text":
        return "⚠ anchors NOT found in loaded guidance text — review catalogue entry"
    return f"catalogue quote, p.{res.get('guidance_page')}"


def _requirement_block(res: dict[str, Any], level: str = "###") -> str:
    lines = [f"{level} {res['guidance_section']} {res['title']} `{res['id']}`", "",
             f"> *Guidance (§{res['guidance_section']}, p.{res['guidance_page']}):* \"{res['quote']}\"", "",
             f"**Status:** {_status_md(res['status'])} · *Grounding:* {_grounding_md(res)}", ""]
    if res.get("applicability_note"):
        lines += [f"*Applicability:* {res['applicability_note']}", ""]
    if res.get("conditional_note"):
        lines += [f"*Note:* {res['conditional_note']}", ""]
    if res["status"] == "not_applicable":
        return "\n".join(lines + ["Not applicable on the basis of the product form / route / production type / additive type entered.", ""])
    if res["status"] == "manual_review":
        lines += ["*Scientific-officer review item: this requirement cannot be assessed by pattern detection.*", ""]
    if res.get("sub_items"):
        lines += ["| Element | Located | Value / evidence | Source |", "|---|---|---|---|"]
        for sub in res["sub_items"]:
            value = sub.get("value", {})
            shown = value.get("value") if value else ""
            src = value.get("citation") if value else ", ".join(sub["citations"][:3])
            qualifier = "" if sub["required"] else " *(where relevant)*"
            if sub.get("note"):
                qualifier += f" *({sub['note']})*"
            lines.append(
                f"| {_md_escape(sub['label'])}{qualifier} | {'yes' if sub['located'] else '**no**'} | "
                f"{_md_escape(shown) or (('“' + _md_escape(sub.get('snippet')) + '”') if sub['located'] else '**[DATA GAP]**')} | {_md_escape(src) or '—'} |"
            )
        lines.append("")
    bc = res.get("batch_check")
    checks = list(res.get("checks", []))
    if bc:
        ids = ", ".join(bc["batch_ids"][:10]) or "none detected"
        checks.insert(0, _check(f"≥ {bc['required_min']} batches", bc["met"],
                                f"{bc['detected']} detected; identifiers: {ids}"))
    if checks:
        lines.append("**Quantitative checks against the guidance:**")
        for c in checks:
            icon = "✅" if c["met"] else ("❌" if c["met"] is False else "❔")
            lines.append(f"- {icon} {c['label']}: {_md_escape(c['detail'])}")
        lines.append("")
    for flag in res.get("flags", []):
        lines += [f"> ⚠ {flag['message']} ({', '.join(flag['citations'])})", ""]
    if res.get("evidence"):
        lines += ["<details><summary>Evidence excerpts</summary>", ""]
        lines += [f"- [{ev['citation']}] {_md_escape(ev['text'])}" for ev in res["evidence"][:5]]
        lines += ["", "</details>", ""]
    if res.get("missing"):
        lines.append("**Data gaps / points to clarify:**")
        lines += [f"- {_md_escape(m)}" for m in res["missing"]]
        lines.append("")
    if res.get("justification"):
        lines.append("**Possible applicant justification located (guidance p.5: reasons should be given for omissions):**")
        lines += [f"- [{j['citation']}] {_md_escape(j['text'])}" for j in res["justification"]]
        lines.append("")
    return "\n".join(lines)


IDENTITY_ROWS = [("generic_name", "Generic name"), ("iupac_name", "Chemical name (IUPAC)"), ("synonyms", "Other names / abbreviations"),
                 ("cas", "CAS No"), ("ec", "EINECS / EC No"), ("molecular_formula", "Molecular formula"),
                 ("structural_formula", "Structural formula"), ("smiles", "openSMILES"), ("molecular_weight", "Molecular weight"),
                 ("flavis", "FLAVIS No"), ("purity", "Purity / specification"), ("water_solubility", "Solubility in water"),
                 ("log_kow", "log Kow"), ("pka", "pKa"), ("melting_point", "Melting point"), ("optical_rotation", "Specific optical rotation")]


def _components_md(result: dict[str, Any]) -> str:
    if not result["components"]:
        return "**[DATA GAP]** No component identified. Enter each active substance/component of the product (§2.1.3, §2.2.1.1)."
    blocks = []
    for i, comp in enumerate(result["components"], start=1):
        c = comp["component"]
        ident = comp["identity"]
        title = c.get("name") or c.get("cas") or f"Component {i}"
        blocks += [f"### Component {i}: {title}", "",
                   f"*Role:* {c.get('role') or 'not stated'}" + (f" · detected from {c['detected_from']}" if c.get("detected_from") else ""), ""]
        if comp.get("cas_check"):
            blocks += [f"*CAS check:* {comp['cas_check']}", ""]
        markers = (result.get("detected", {}).get("all", {}) or {}).get("markers", [])
        if "plant origin" in (c.get("role") or "") and markers:
            blocks += ["| Constituent / marker compound (§2.1.3, §2.2.1.1) | Specification | Source |", "|---|---|---|"]
            blocks += [f"| {_md_escape(m['label'])} | {_md_escape(m['specification'])} | {m['citation']} |" for m in markers]
            blocks.append("")
        blocks += ["| Identity element (§2.2.1.1 / §2.2.2.1) | Value as stated in the dossier | Source |", "|---|---|---|"]
        for key, label in IDENTITY_ROWS:
            v = ident.get(key)
            blocks.append(f"| {label} | {_md_escape(v['value']) if v else '**[not located]**'} | {v['citation'] if v else '—'} |")
        blocks.append("")
        for res in comp["requirements"]:
            blocks.append(_requirement_block(res, level="####"))
    return "\n".join(blocks)


def _gap_table_md(result: dict[str, Any]) -> str:
    rows = ["| § | Requirement | Scope | Status | Missing / to clarify | Justification located |", "|---|---|---|---|---|---|"]
    order = {"gap": 0, "partial": 1, "conditional": 2, "manual_review": 3, "evidence_located": 4, "not_applicable": 5}
    entries = [(res, "Additive") for res in result["requirements"]]
    for comp in result["components"]:
        name = comp["component"].get("name") or comp["component"].get("cas") or "component"
        entries += [(res, name) for res in comp["requirements"]]
    entries.sort(key=lambda e: (order.get(e[0]["status"], 9), [int(x) for x in e[0]["guidance_section"].split(".")]))
    for res, scope in entries:
        missing = "; ".join(res.get("missing", [])[:4]) or "—"
        just = ", ".join(j["citation"] for j in res.get("justification", [])) or "—"
        rows.append(f"| {res['guidance_section']} | {_md_escape(res['title'])} `{res['id']}` | {_md_escape(scope)} | "
                    f"{_status_md(res['status'])} | {_md_escape(missing)} | {just} |")
    s = result["summary"]
    head = (f"Not located: **{s.get('gap', 0)}** · Partially addressed: **{s.get('partial', 0)}** · "
            f"Applicability to confirm: **{s.get('conditional', 0)}** · SO review: **{s.get('manual_review', 0)}** · "
            f"Evidence located: **{s.get('evidence_located', 0)}** · Not applicable: **{s.get('not_applicable', 0)}**")
    flags = []
    for f in result.get("scope_flags", []):
        cites = ", ".join(e["citation"] for e in f.get("evidence", [])[:3])
        flags.append(f"> **{f['flag']}.** {f['detail']}{(' Evidence: ' + cites + '.') if cites else ''}")
    general = ("Per the guidance (p.5): *\"Reasons should be given for the omission from the dossier of any data prescribed there.\"* "
               "For each gap below, check whether the applicant has justified the omission.")
    return "\n".join([head, "", general, ""] + sum(([f, ""] for f in flags), []) + rows)


def _documents_md(result: dict[str, Any]) -> str:
    if not result["documents"]:
        return "**No documents were processed.**"
    rows = ["| ID | Document | Type (auto-classified) | Pages | Characters | Warnings |", "|---|---|---|---|---|---|"]
    for d in result["documents"]:
        rows.append(f"| {d['doc_id']} | {_md_escape(d['name'])} | {_md_escape(d['doc_type'])} | {d['pages']} | {d['characters']} | {_md_escape('; '.join(d['warnings'])) or '—'} |")
    notes = [f"- {_md_escape(n)}" for n in result.get("fetch_notes", [])]
    return "\n".join(rows + ([""] + ["**Retrieval notes:**"] + notes if notes else []))


def _grounding_section_md(result: dict[str, Any], catalogue: dict[str, Any]) -> str:
    g = catalogue["guidance"]
    lines = [f"- Guidance: {g['citation']} [doi:{g['doi']}]({g['url']}). {g['license']}.",
             f"- Legal basis: {g['legal_basis']}.",
             f"- Requirement catalogue `{catalogue['catalogue_id']}` v{catalogue['catalogue_version']} — every requirement quotes the guidance "
             "verbatim with section and page. Scope: " + catalogue["scope"]]
    if result["guidance_text_loaded"]:
        lines.append("- Guidance text loaded: each requirement was re-anchored against the guidance text at run time "
                     "(grounding column). Requirements marked ⚠ could not be anchored and must not be relied upon.")
    else:
        lines.append("- Guidance text not loaded in this instance: grounding relies on the catalogue quotes (verified against "
                     "the guidance by the automated test suite).")
    lines.append("- Extraction is deterministic (pattern-based); every value carries a document/page citation. "
                 "Statuses indicate whether evidence was *located*, not whether it is adequate or compliant.")
    return "\n".join(lines)


def render_markdown(result: dict[str, Any], catalogue: dict[str, Any], template_path: Optional[Path] = None) -> str:
    if template_path is None:
        template_path = LOCAL_TEMPLATE_PATH if LOCAL_TEMPLATE_PATH.exists() else TEMPLATE_PATH
    template = re.sub(r"<!--.*?-->\s*", "", template_path.read_text(encoding="utf-8"), flags=re.S)  # authoring notes
    meta = result["meta"]
    by_section: dict[str, list[str]] = {}
    for res in result["requirements"]:
        by_section.setdefault(res["section"], []).append(_requirement_block(res))
    product = meta.get("product_name") or "[product name not entered]"
    header = "\n".join([
        *([f"*Applicant:* {meta['applicant']} · *Dossier / question reference:* {meta.get('dossier_ref') or '—'}  "]
          if meta.get("applicant") or meta.get("dossier_ref") else []),
        f"*Product form:* {meta.get('product_form') or 'not stated'} · *Routes:* {', '.join(meta.get('routes') or []) or 'not stated'} · "
        f"*Production type:* {', '.join(meta.get('production_types') or []) or 'not stated'} · *Additive type:* {meta.get('additive_type') or 'not stated'}  ",
        f"*Generated:* {result['generated_at']} by KEvidence step 0 (chemical characterisation MVP) against "
        f"{catalogue['guidance']['citation']}  ",
        "**Draft for scientific-officer verification — contains confidential dossier information.**",
    ])
    caveats = ("---\n*This draft structures the applicant's evidence against the FEEDAP (2017) guidance on identity, characterisation "
               "and conditions of use of feed additives. It is not an EFSA assessment or conclusion. Values are reproduced as stated "
               "in the dossier and must be verified against the source documents. A 'not located' element can result from "
               "extraction limits (scanned PDFs, images, unusual table layouts) as well as from a true data gap.*")
    narrative = result.get("narrative") or {}
    replacements = {
        "PRODUCT_NAME": product,
        "HEADER": header,
        "DOCUMENT_REGISTER": _documents_md(result),
        "GAP_ANALYSIS": _gap_table_md(result),
        "GROUNDING": _grounding_section_md(result, catalogue),
        "CAVEATS": caveats,
        "OPINION_111": _opinion_111_md(result),
        "DETECTED": _detected_md(result),
        "APPENDIX_A": _appendix_a_md(result),
        "OPINION_112": _opinion_112_md(result),
    }
    for section in catalogue["template_sections"]:
        token = section["token"]
        replacements[token] = "\n".join(by_section.get(token, []))
    replacements["S2_2"] = "\n".join(by_section.get("S2_2", []) + [_components_md(result)])
    for token, text in narrative.items():
        if token in replacements and text.get("text"):
            flag = ("all numbers traced to dossier excerpts" if text.get("grounded")
                    else f"⚠ values not traced to dossier excerpts: {', '.join(text.get('untraced', []))}")
            replacements[token] = (f"**Draft narrative (LLM-assisted; {flag})**\n\n> {text['text']}\n\n" + replacements[token])

    def sub(m):
        return replacements.get(m.group(1), m.group(0))
    return re.sub(r"\{\{([A-Z0-9_]+)\}\}", sub, template)


# ---------------------------------------------------------------------------
# Optional LLM drafting (off by default)
# ---------------------------------------------------------------------------

NARRATIVE_SYSTEM_PROMPT = """You draft the characterisation section of an EFSA FEEDAP scientific opinion.
Rules (strict):
1. Use ONLY the evidence excerpts provided. Never add values, batch numbers, methods or conclusions that are not in the excerpts.
2. Follow the requirement text, which comes from the FEEDAP (2017) identity/characterisation guidance (EFSA Journal 15(10):5023).
3. Report values as stated by the applicant ("The applicant reported…"), cite the source in square brackets, e.g. [D2 p.4].
4. Where the evidence is missing or below the guidance minimum (e.g. fewer batches than required), say so explicitly.
5. Do not express safety or efficacy conclusions.
Return JSON: {"sections": {"<SECTION_TOKEN>": "<one paragraph>"}} for the section tokens supplied."""


def _numbers(text: str) -> set[str]:
    return {n.rstrip(".").replace(",", ".") for n in re.findall(r"\d+(?:[.,]\d+)?", text or "")}


def draft_narrative(result: dict[str, Any], complete: Callable[[str, str], str], max_chars: int = 24000) -> dict[str, Any]:
    """Draft one paragraph per section with an LLM and check every number against the evidence."""
    sections: dict[str, list[dict[str, Any]]] = {}
    for res in result["requirements"]:
        sections.setdefault(res["section"], []).append(res)
    for comp in result["components"]:
        for res in comp["requirements"]:
            sections.setdefault("S2_2", []).append(res)
    payload, evidence_text = [], []
    for token, items in sections.items():
        block = {"section": token, "requirements": []}
        for res in items:
            ev = [f"[{e['citation']}] {e['text']}" for e in res.get("evidence", [])[:6]]
            evidence_text += ev
            block["requirements"].append({
                "id": res["id"], "guidance_section": res["guidance_section"], "requirement": res["requirement"], "status": res["status"],
                "missing": res.get("missing", []), "batch_check": res.get("batch_check"), "evidence": ev,
            })
        payload.append(block)
    user = json.dumps(payload)[:max_chars]
    raw = complete(NARRATIVE_SYSTEM_PROMPT, user)
    try:
        parsed = json.loads(re.search(r"\{.*\}", raw, re.S).group(0))
        drafted = parsed.get("sections", {})
    except Exception:
        return {}
    allowed = _numbers(" ".join(evidence_text) + " " + user)
    out = {}
    for token, text in drafted.items():
        if not isinstance(text, str):
            continue
        citation_numbers = _numbers(" ".join(re.findall(r"\[[^\]]*\]", text)))
        untraced = sorted(n for n in _numbers(text) - allowed - citation_numbers)
        out[token] = {"text": text.strip(), "grounded": not untraced, "untraced": untraced}
    return out


# ---------------------------------------------------------------------------
# Assistant grounding helper
# ---------------------------------------------------------------------------

QUESTION_STOPWORDS = {
    "what", "which", "when", "where", "does", "should", "would", "could", "must", "with", "from", "that", "this",
    "there", "their", "they", "them", "have", "about", "into", "than", "then", "also", "only", "each", "other",
    "guidance", "require", "required", "requires", "requirement", "requirements", "explain", "tell", "please",
    "need", "needed", "needs", "data", "information", "provide", "provided", "additive", "additives", "feed",
    "efsa", "feedap", "dossier", "applicant", "says", "much", "many", "matter", "most", "exactly", "gaps",
}


def guidance_context_for_question(question: str, catalogue: Optional[dict[str, Any]] = None,
                                  guidance_pages: Optional[list[str]] = None, limit: int = 6) -> tuple[str, list[str]]:
    """Select catalogue requirements (and guidance passages, when loaded) relevant to a question."""
    catalogue = catalogue or load_catalogue()
    words = {w for w in re.findall(r"[a-z]{4,}", question.lower())} - QUESTION_STOPWORDS
    scored = []
    for req in catalogue["requirements"]:
        hay = f"{req['title']} {req['requirement']} {req['quote']} {' '.join(req.get('guidance_anchors', []))}".lower()
        score = sum(1 for w in words if w in hay)
        if re.search(re.escape(req["id"]), question, re.I) or re.search(r"(§|section\s*)" + re.escape(req["guidance_section"]) + r"\b", question, re.I):
            score += 10
        if score:
            scored.append((score, req))
    scored.sort(key=lambda x: -x[0])
    chosen = [r for _, r in scored[:limit]]
    anchors = anchor_requirements(catalogue, guidance_pages)
    parts = []
    for req in chosen:
        g = anchors[req["id"]]
        part = (f"[{req['id']}] Guidance §{req['guidance_section']} (p.{req['guidance_page']}): {req['title']}\n"
                f"Requirement (paraphrase): {req['requirement']}\nVerbatim guidance quote: \"{req['quote']}\"\n"
                f"Grounding status: {g['status']}")
        for p in g.get("passages", [])[:2]:
            part += f"\nGuidance text p.{p['page']}: \"{p['excerpt']}\""
        parts.append(part)
    return "\n\n".join(parts), [r["id"] for r in chosen]
