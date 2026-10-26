"""Word (.docx) export of the step 0 characterisation.

The on-screen EFSA-layout output is Markdown (characterisation.render_markdown). This module turns exactly that
Markdown into a Word document, so the download and the screen cannot disagree. It handles the subset the renderer
produces: headings, paragraphs with hard line breaks, pipe tables, bullet lists, blockquotes, horizontal rules,
bold/italic/code, and footnotes ([^id] references with [^id]: definitions).

Footnotes are written as numbered superscript markers with the note text in a small paragraph where the definition
stands in the Markdown (directly under its table), as the opinion layout places table notes.
"""

from __future__ import annotations

import datetime as _dt
import io
import re
from typing import Any, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MAX_MARKDOWN_CHARS = 3_000_000
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

BODY_FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
HEADER_FILL = "D9E2F3"
QUOTE_FILL = "F2F2F2"

_INLINE = re.compile(
    r"(?P<bold>\*\*(?P<b>.+?)\*\*)"
    r"|(?P<italic>(?<![\w*])\*(?P<i>[^*\s](?:[^*]*[^*\s])?)\*(?![\w*]))"
    r"|(?P<code>`(?P<c>[^`]+)`)"
    r"|(?P<fn>\[\^(?P<fid>[^\]\s]+)\])"
)
_ESCAPED = re.compile(r"\\([\\`*_{}\[\]()#+\-.!|<>])")
_TABLE_SEP = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_FN_DEF = re.compile(r"^\[\^(?P<id>[^\]\s]+)\]:\s*(?P<text>.*)$")


def _unescape(text: str) -> str:
    return _ESCAPED.sub(lambda m: m.group(1), text)


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


class _Footnotes:
    """Numbers footnotes in order of first reference."""

    def __init__(self) -> None:
        self.numbers: dict[str, int] = {}

    def number(self, fid: str) -> int:
        if fid not in self.numbers:
            self.numbers[fid] = len(self.numbers) + 1
        return self.numbers[fid]


def _set_cell_fill(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _shade_paragraph(paragraph: Any, fill: str, left_border: Optional[str] = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if left_border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), left_border)
        borders.append(left)
        p_pr.append(borders)


def _bottom_rule(paragraph: Any) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")
    borders.append(bottom)
    p_pr.append(borders)


def _add_field(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instruction:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instruction} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _add_inline(paragraph: Any, text: str, notes: _Footnotes, size: Optional[float] = None, bold: bool = False,
                italic: bool = False) -> None:
    """Add Markdown-inline text to a paragraph."""
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            _run(paragraph, _unescape(text[pos:m.start()]), size, bold, italic)
        if m.group("bold"):
            _add_inline(paragraph, m.group("b"), notes, size, True, italic)
        elif m.group("italic"):
            _add_inline(paragraph, m.group("i"), notes, size, bold, True)
        elif m.group("code"):
            r = _run(paragraph, m.group("c"), size, bold, italic)
            r.font.name = "Consolas"
        else:
            r = _run(paragraph, str(notes.number(m.group("fid"))), size, bold, italic)
            r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        _run(paragraph, _unescape(text[pos:]), size, bold, italic)


def _run(paragraph: Any, text: str, size: Optional[float], bold: bool, italic: bool) -> Any:
    r = paragraph.add_run(text)
    if size:
        r.font.size = Pt(size)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    return r


def _add_paragraph_lines(doc: Any, lines: list[str], notes: _Footnotes, style: Optional[str] = None, size: Optional[float] = None) -> Any:
    """A paragraph from consecutive Markdown lines; a line ending in two spaces ends in a hard line break."""
    p = doc.add_paragraph(style=style)
    for i, raw in enumerate(lines):
        hard = raw.endswith("  ")
        _add_inline(p, raw.strip(), notes, size)
        if i < len(lines) - 1:
            if hard:
                p.add_run().add_break(WD_BREAK.LINE)
            else:
                p.add_run(" ")
    return p


TEXT_WIDTH_CM = 17.0


def _column_widths(rows: list[list[str]], ncols: int) -> list[Any]:
    """Column widths from content length (square-root damped, with a floor), filling the text width."""
    weights = []
    for ci in range(ncols):
        lengths = [len(_unescape(r[ci])) for r in rows if ci < len(r)] or [1]
        weights.append(max(4.0, (sum(lengths) / len(lengths)) ** 0.5 * 2 + max(lengths) ** 0.5))
    total = sum(weights)
    # no column narrower than its longest header word (e.g. "Justification"), so headings never break mid-word
    floors = [max(1.4, 0.2 * max((len(w) for w in _unescape(rows[0][ci]).split()), default=0) + 0.3) if ci < len(rows[0]) else 1.4
              for ci in range(ncols)]
    return [Cm(max(floors[ci], TEXT_WIDTH_CM * weights[ci] / total)) for ci in range(ncols)]


def _add_table(doc: Any, rows: list[list[str]], notes: _Footnotes) -> None:
    ncols = max(len(r) for r in rows)
    size = 7.5 if ncols >= 7 else 8.5 if ncols >= 5 else 9
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = _column_widths(rows, ncols)
    table.autofit = False
    for ci, w in enumerate(widths):
        table.columns[ci].width = w      # the grid widths; Word and LibreOffice both honour these
    for ri, row in enumerate(rows):
        cells = table.rows[ri].cells
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.width = widths[ci]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            _add_inline(p, text, notes, size, bold=(ri == 0))
            if ri == 0:
                _set_cell_fill(cell, HEADER_FILL)
        if ri == 0:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            el = OxmlElement("w:tblHeader")
            el.set(qn("w:val"), "true")
            tr_pr.append(el)
        cant_split = OxmlElement("w:cantSplit")
        table.rows[ri]._tr.get_or_add_trPr().append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _configure(doc: Any, title: str, header_text: str) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, before in (("Title", 20, 0), ("Heading 1", 16, 18), ("Heading 2", 13, 14), ("Heading 3", 11, 10), ("Heading 4", 10, 8)):
        st = styles[name]
        st.font.name = BODY_FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        st.element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        st.element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        styles[name].font.name = BODY_FONT
        styles[name].paragraph_format.space_after = Pt(2)

    header = sec.header.paragraphs[0]
    header.text = ""
    r = header.add_run(header_text)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    footer = sec.footer.paragraphs[0]
    footer.text = ""
    for text, field in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
        fr = footer.add_run(text)
        fr.font.size = Pt(8)
        _add_field(footer, field)
    footer.alignment = 2  # right

    cp = doc.core_properties
    cp.title = title
    cp.author = "KEvidence"
    cp.subject = "Step 0 characterisation, draft for scientific-officer verification"
    cp.created = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None)


def markdown_to_docx(markdown: str, title: str = "Characterisation of the additive", header_text: Optional[str] = None) -> bytes:
    """Convert the step 0 Markdown to a .docx and return its bytes."""
    if len(markdown) > MAX_MARKDOWN_CHARS:
        raise ValueError("Markdown is too large to export")
    doc = Document()
    _configure(doc, title, header_text or "KEvidence · draft for scientific-officer verification · may contain confidential dossier information")
    notes = _Footnotes()
    lines = markdown.replace("\r\n", "\n").split("\n")
    i, n, first_h1 = 0, len(lines), True
    para: list[str] = []

    def flush() -> None:
        if para:
            _add_paragraph_lines(doc, para, notes)
            para.clear()

    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush()
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            level = len(m.group(1))
            if level == 1 and first_h1:
                p = doc.add_paragraph(style="Title")
                first_h1 = False
            else:
                if level == 1:
                    doc.add_page_break()
                p = doc.add_paragraph(style=f"Heading {min(level, 4)}")
            _add_inline(p, m.group(2).strip(), notes)
            i += 1
            continue
        if re.match(r"^\s*(-{3,}|\*{3,})\s*$", line):
            flush()
            _bottom_rule(doc.add_paragraph())
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < n and _TABLE_SEP.match(lines[i + 1].strip()):
            flush()
            rows = [_split_row(line)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            _add_table(doc, rows, notes)
            continue
        fm = _FN_DEF.match(stripped)
        if fm:
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            num = _run(p, f"{notes.number(fm.group('id'))} ", 8, False, False)
            num.font.superscript = True
            _add_inline(p, fm.group("text"), notes, 8)
            i += 1
            continue
        bm = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bm:
            flush()
            p = doc.add_paragraph(style="List Bullet 2" if len(bm.group(1)) >= 2 else "List Bullet")
            _add_inline(p, bm.group(2), notes)
            i += 1
            continue
        nm = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if nm:
            flush()
            _add_inline(doc.add_paragraph(style="List Number"), nm.group(1), notes)
            i += 1
            continue
        if stripped.startswith(">"):
            flush()
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            p = _add_paragraph_lines(doc, quote, notes)
            p.paragraph_format.left_indent = Cm(0.5)
            _shade_paragraph(p, QUOTE_FILL, left_border="7F7F7F")
            continue
        para.append(line)
        i += 1
    flush()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
